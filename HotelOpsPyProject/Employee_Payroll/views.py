from itertools import groupby
from operator import attrgetter
from django.shortcuts import render,redirect
from Employee_Payroll.forms import AlifFileUploadForm
from hotelopsmgmtpy.GlobalConfig import MasterAttribute,OrganizationDetail
from app.models import OrganizationMaster,EmployeeMaster
from .models import AlifCSVPunchRecord, Attendance_Data,Update_Attendance_Request,Update_Attendance_Request_V1,SalarySlip,Raw_Attendance_Data,Organization_Details,SalaryEmails,Raw_Attendance_Data_File,WeekOffDetails,PayrollErrorLog,AttendanceLock,ShfitMaster, AlifCSVPunchRecord_Log
from Leave_Management_System.models import  Leave_Type_Master,Emp_Leave_Balance_Master
import calendar
import requests
from django.contrib import messages
from django.db import   transaction
from datetime import datetime, timedelta
from datetime import datetime
from Employee_Payroll.models import Raw_Attendance_Data, Attendance_Data, SalaryAttendance
from Leave_Management_System.models import Leave_Application
from django.db.models import Q

from django.db.models import Min, Max
from datetime import datetime
import csv
import pandas as pd
from django.shortcuts import get_object_or_404


from django.db import transaction
from django.utils import timezone


from django.http import HttpResponse
from io import BytesIO
from django.template.loader import get_template
from xhtml2pdf import pisa
from app.views import EmployeeDataSelect,EmployeeDataSelectForSalary


    

@transaction.atomic
def Upload_Attendace(request):
    if 'OrganizationID' not in request.session:
        return redirect(MasterAttribute.Host)
    else:
        print("Show Page Session")

    OrganizationID = request.session["OrganizationID"]
   
    UserID = str(request.session["UserID"])
    UserType = request.session["UserType"]
    
    # emp_list =  EmployeeMaster.objects.filter(OrganizationID=OrganizationID,IsDelete=False)
    emp_list =  EmployeeDataSelect(OrganizationID)
    
     
    try :
        org_obj  = Organization_Details.objects.get(OID=OrganizationID,IsDelete=False)
        OID_Code  = org_obj.OID_Code
        UploadFormatType = org_obj.UploadFormatType
        DownloadFormat = org_obj.DownloadFormat.url
        
    except Organization_Details.DoesNotExist:
        
        OID_Code = ''
        UploadFormatType = '.txt'
        DownloadFormat =  '/Employee_Payroll/DownloadFormat/Defaultformat.txt'
      

    
    
    with transaction.atomic():
        if request.method == "POST":
            file = request.FILES.get('file')
            Attendance_Date = request.POST['Attendance_Date']            
            attendance_date_obj = datetime.strptime(Attendance_Date, '%Y-%m-%d').date()
            
            parsed_date = datetime.strptime(Attendance_Date, '%Y-%m-%d')
            messages_Date = parsed_date.strftime('%d-%m-%Y')
            
          
            if UploadFormatType != "." +Path(file.name).suffix.lower()[1:]:
                    messages.warning(request,f'Upload format type is not macthed to  {UploadFormatType}')
                    return redirect('Upload_Attendace')
            
              
            
            try:
                Attendance_check = Raw_Attendance_Data_File.objects.get(Attendance_Date = Attendance_Date,OrganizationID = OrganizationID,IsDelete = False)
                messages.warning(request, f"Attendance record is present for {messages_Date}")    
                return redirect('Upload_Attendace')
            except Raw_Attendance_Data_File.DoesNotExist:    
                decoded_file = file.read().decode('utf-8').splitlines()
                Attendance_create = Raw_Attendance_Data_File.objects.create(File_Name=file.name, Attendance_Date=Attendance_Date, OrganizationID=OrganizationID, CreatedBy=UserID)

                match_found = False 
                    
                for row in decoded_file:
                    fields = row.split(',')
                    if len(fields) == 4:  
                        RawEmployeeCode, Date, Time, Status = fields
                        date_obj = datetime.strptime(Date, '%d/%m/%Y').date()
                        DumpEmployeeCode = RawEmployeeCode
                        if OID_Code !='':
                                if len(RawEmployeeCode) == 1:
                                        RawEmployeeCode = '00' + RawEmployeeCode
                                elif len(RawEmployeeCode) == 2:
                                        RawEmployeeCode = '0' + RawEmployeeCode

                                    
                                DumpEmployeeCode = str(OID_Code) + RawEmployeeCode

                                
                                DumpEmployeeCode = DumpEmployeeCode.zfill(9)
                        
                        if date_obj == attendance_date_obj: 
                            raw_attendance = Raw_Attendance_Data.objects.create(
                                EmployeeCode = DumpEmployeeCode,
                                Date=date_obj, 
                                Time=Time,  
                                Status=Status,
                                OrganizationID=OrganizationID,
                                CreatedBy=UserID
                            )
                            match_found = True  
                            

                if not match_found:  
                    Attendance_create.IsDelete = True
                    Attendance_create.save() 
                    messages.warning(request, f"Attendance record of file does not match with selected date {messages_Date}")    
                    return redirect('Upload_Attendace')

   
                for emp in emp_list:
                    Level = emp.Level
                    Level_List = ["M","M1","M2","M3","M4","M5","M6"]
                   
                    EmployeeCode = emp['EmployeeCode']
                    
                   
                    
                    latest_out_time_query = f"""
                        SELECT MAX(Time) AS MaxOutTime
                        FROM Employee_Payroll_raw_attendance_data
                        WHERE EmployeeCode = '{EmployeeCode}'
                            AND OrganizationID = {OrganizationID}
                            AND IsDelete = 0
                            AND Status = 'OUT'
                            AND Date >= '{attendance_date_obj}'
                            AND Date < '{attendance_date_obj + timedelta(days=1)}';
                    """
                    earliest_in_time_query = f"""
                        SELECT MIN(Time) AS MinInTime
                        FROM Employee_Payroll_raw_attendance_data
                        WHERE EmployeeCode = '{EmployeeCode}'
                            AND Date = '{attendance_date_obj}'
                            AND OrganizationID = {OrganizationID}
                            AND IsDelete = 0
                            AND Status = 'IN';
                    """
                    with connection.cursor() as cursor:
                        cursor.execute(latest_out_time_query)
                        latest_out_time_result = cursor.fetchone()

                        cursor.execute(earliest_in_time_query)
                        earliest_in_time_result = cursor.fetchone()

                    latest_out_time = latest_out_time_result[0] if latest_out_time_result else None
                    earliest_in_time = earliest_in_time_result[0] if earliest_in_time_result else None
                    if (latest_out_time and earliest_in_time) or (Level in Level_List and (latest_out_time or earliest_in_time)):
                        attendance_records = {
                            'EmployeeCode': EmployeeCode,
                            'Date': attendance_date_obj,
                            'MinInTime': earliest_in_time,
                            'MaxOutTime': latest_out_time
                        }
                        in_time_str = attendance_records['MinInTime']
                        out_time_str = attendance_records['MaxOutTime']
                   
                        
                        if Level in Level_List:
                               
                                if in_time_str and out_time_str:
                                    in_time = datetime.strptime(in_time_str, '%H:%M:%S')
                                    out_time = datetime.strptime(out_time_str, '%H:%M:%S')
                                    duty_hours = out_time - in_time
                                    duty_hours_time = str(duty_hours).split()[-1]
                                    attendance_records['DutyHours'] = duty_hours_time

                                    duty_hours_timedelta = datetime.strptime(duty_hours_time, '%H:%M:%S') - datetime(1900, 1, 1)
                                else:
                                    attendance_records['DutyHours'] = None
                                    duty_hours_timedelta = timedelta(hours=0)
                              
                                status = "Absent"
                                if in_time_str or out_time_str:
                                        status = "Present"
                                attendance_records['status'] = status

                        else:   
                               
                                     
                                if in_time_str and out_time_str:
                                    in_time = datetime.strptime(in_time_str, '%H:%M:%S')
                                    out_time = datetime.strptime(out_time_str, '%H:%M:%S')
                                    duty_hours = out_time - in_time
                                    duty_hours_time = str(duty_hours).split()[-1]
                                    attendance_records['DutyHours'] = duty_hours_time

                                    duty_hours_timedelta = datetime.strptime(duty_hours_time, '%H:%M:%S') - datetime(1900, 1, 1)
                                else:
                                    attendance_records['DutyHours'] = None
                                    duty_hours_timedelta = timedelta(hours=0)

                                status = "Absent"
                                if duty_hours_timedelta is not None:
                                    if duty_hours_timedelta >= timedelta(hours=8, minutes=30):
                                        status = "Present"
                                    elif duty_hours_timedelta >= timedelta(hours=5):
                                        status = "Half Day Present"
                                attendance_records['status'] = status
                        
                        if attendance_records:
                            Attendance_Data.objects.create(
                                EmployeeCode=EmployeeCode,
                                Date=Attendance_Date,
                                In_Time=attendance_records['MinInTime'],
                                Out_Time=attendance_records['MaxOutTime'],
                                Duty_Hour=attendance_records['DutyHours'],
                                OrganizationID=OrganizationID,
                                CreatedBy=UserID,
                                Status=attendance_records['status'],
                                IsUpload=True
                            )
                        
                    else:
                       
                        MinInTime = None
                        MaxOutTime = None
                        Duty_Hour = None
                        try:
                            
                           
                            
                            query = """
                            SELECT * 
                            FROM Leave_Management_System_leave_application 
                            WHERE Start_Date <= %s
                            AND End_Date >= %s
                            AND IsDelete = 0
                            AND Emp_code = %s
                            AND OrganizationID = %s
                            AND Status = 1;
                            """ 
                            with connection.cursor() as cursor:
                                cursor.execute(query, [Attendance_Date,Attendance_Date,EmployeeCode, OrganizationID])
                                rows = cursor.fetchall()
                                columns = [col[0] for col in cursor.description]
                            rowslist = [dict(zip(columns, row)) for row in rows] 
                            
                            if rowslist:
                                leave_id = rowslist[0].get('Leave_Type_Master_id') 
                                if leave_id is not None:
                                    leave = get_object_or_404(Leave_Type_Master, id=leave_id, IsDelete=False,Is_Active=True)
                                    status = leave.Type

                                attendance_record = Attendance_Data.objects.create(EmployeeCode=EmployeeCode, Date=Attendance_Date, In_Time=MinInTime, Out_Time=MaxOutTime, OrganizationID=OrganizationID, Status=status, IsUpload=True, CreatedBy=UserID, Duty_Hour=Duty_Hour)
                       
                        except Leave_Application.DoesNotExist:
                                status = "Absent"
                                attendance_record = Attendance_Data.objects.create(EmployeeCode=EmployeeCode, Date=Attendance_Date, In_Time=MinInTime, Out_Time=MaxOutTime, OrganizationID=OrganizationID, Status=status, IsUpload=True, Duty_Hour=Duty_Hour, CreatedBy=UserID)
 
                                

                                
                

                messages.success(request, f"Uploaded Successfully for {messages_Date}")    
                return redirect('/Employee_Payroll/Daily_Attendace/?Q_Date='+str(Attendance_Date))

    context = {'UploadFormatType':UploadFormatType,'DownloadFormat' : DownloadFormat}
    return render(request, "EMP_PAY/Attendance/Upload_Attendace.html", context)


import pandas as pd
from io import BytesIO
from django.http import HttpResponse,Http404
from django.core.files.storage import FileSystemStorage
from datetime import datetime

import pandas as pd
from io import BytesIO
from django.http import HttpResponse
from django.core.files.storage import FileSystemStorage
from datetime import datetime
from django.shortcuts import redirect

import pandas as pd
from io import BytesIO
from django.http import HttpResponse
from django.core.files.storage import FileSystemStorage
from datetime import datetime
from django.shortcuts import redirect



from .azure import upload_file_to_blob,download_blob
from .models import AttendanceSalaryFile

@transaction.atomic
def UploadAttendanceSalaryFile(request):
    if 'OrganizationID' not in request.session:
        return redirect(MasterAttribute.Host)
    
    OrganizationID = request.session["OrganizationID"]
    UserID = str(request.session["UserID"])
    
    if request.method == 'POST':
        file = request.FILES.get('doc')  # Excel file
        pdf_file = None
        if 'pdf' in request.FILES:
            pdf_file = request.FILES.get('pdf')  # PDF file

        file_type = request.POST.get('Type')
        month = request.POST.get('month')
        year = request.POST.get('year')

        # Check if a previous record exists
        Previousattendance_salary_file = AttendanceSalaryFile.objects.filter(
            Type=file_type,
            Month=month,
            Year=year,
            OrganizationID=OrganizationID,IsDelete=False
        )
        for Pr in Previousattendance_salary_file:
        
            if Pr:
                Pr.IsDelete = True
                Pr.ModifyBy = UserID
                Pr.save()

        file_title = f'{file_type}_{month}_{year}'

        # Create a new AttendanceSalaryFile instance
        attendance_salary_file = AttendanceSalaryFile.objects.create(
            Type=file_type,
            Month=month,
            Year=year,
            FileTitle=file_title,
            PdfFileTitle=file_title if pdf_file is not None else "",

            OrganizationID=OrganizationID,
            CreatedBy=UserID,
        )
        
        id = attendance_salary_file.id
        
        # Upload Excel file to Azure Blob Storage
        if file:
            upload_file_to_blob(file, id)  # For Excel file

        # Upload PDF file to Azure Blob Storage
        if pdf_file is not None:
            pdf_id = id  # Using the same ID for the PDF upload; you can adjust if needed
            upload_file_to_blob(pdf_file, pdf_id, is_pdf=True)  # Set is_pdf to True for PDF file
        
        return redirect(reverse('AttendanceSalaryList') + f'?Type={file_type}&month={month}&year={year}')




import mimetypes
import os  # Import os for handling file names
from django.http import HttpResponse, Http404
from django.contrib import messages

import os
import mimetypes
from django.http import HttpResponse, Http404
from django.contrib import messages
from .models import AttendanceSalaryFile  

def download_file(request):
    if 'OrganizationID' not in request.session:
        return redirect(MasterAttribute.Host)

    OrganizationID = request.session["OrganizationID"]
    UserID = str(request.session["UserID"])
    DomainCode = request.session["DomainCode"]
    id = request.GET.get('DID')
    file_type = request.GET.get('type') 
    
    file = AttendanceSalaryFile.objects.filter(id=id).first()
    if not file:
        raise Http404("File not found")

    if file_type == 'excel' and file.FileName:
        FileTitle = file.FileTitle  
        FileName = file.FileName
    elif file_type == 'pdf' and file.PdfFileName:
        FileTitle = file.FileTitle  
        FileName = file.PdfFileName
    else:
        raise Http404("File not found")

    _, suffix = os.path.splitext(FileName)

    file_type, _ = mimetypes.guess_type(FileName)
    if file_type is None:
        file_type = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet' if suffix in ['.xls', '.xlsx'] else 'application/pdf'

    print(f"Guessing MIME type for '{FileName}': {file_type}")

    blob_content = download_blob(FileName)

    if blob_content:
        response = HttpResponse(blob_content, content_type=file_type)
        response['Content-Disposition'] = f'attachment; filename="{DomainCode}_{FileTitle}{suffix}"'
      
        return response

    raise Http404("File not found")



def get_current_month_abbr():
    months = ['Jan', 'Feb', 'Mar', 'Apr', 'May', 'Jun', 'Jul', 'Aug', 'Sep', 'Oct', 'Nov', 'Dec']
    
    current_month = datetime.now().month
    
    return months[current_month - 1]


@transaction.atomic
def AttendanceSalaryList(request):
    if 'OrganizationID' not in request.session:
        return redirect(MasterAttribute.Host)

    OrganizationID = request.session["OrganizationID"]
    UserID = str(request.session["UserID"])
    Type = request.GET.get('Type','Attendance')

    today = datetime.now()
    
    if 'year' in request.GET:
        year = int(request.GET.get('year'))
    else:
        year = today.year  

    if 'month' in request.GET:
        month = request.GET.get('month')
    else:
        month = get_current_month_abbr()  
    year  = year
    month = month
    if month == 'All':
        attsalobj = AttendanceSalaryFile.objects.filter(
                    IsDelete=False,
                    OrganizationID=OrganizationID,
                    Type=Type,
                  
                    Year=year
                )
    else:
         attsalobj = AttendanceSalaryFile.objects.filter(
                    IsDelete=False,
                    OrganizationID=OrganizationID,
                    Type=Type,
                     Month=month,
                    Year=year
                )    
    

    context = {
        'attsalobj': attsalobj,
        'CYear': range(today.year, 2020, -1),
        'CMonth': today.month,
        'month': month,
        'Type':Type
        
    } 
    return render(request, "EMP_PAY/Attendance/AttendanceSalaryList.html", context)


def AttendaceCorrectEtmExlUp(request):
    if 'OrganizationID' not in request.session:
        return redirect(MasterAttribute.Host)
    
    OrganizationID = request.session["OrganizationID"]
    UserID = str(request.session["UserID"])
    EmployeeCode = request.session["EmployeeCode"]

    if request.method == "POST":
        print("hello POST")
        file = request.FILES['doc']
        fs = FileSystemStorage()
        filename = fs.save(file.name, file)
        file_path = fs.path(filename)

        df = pd.read_excel(file_path)

        attendance_data = []

        employee_code_col = 1  

      
        row_index = 8  
        from_date_str =df.iloc[2, 0]
        print("FromDate",from_date_str)
        to_date_str = df.iloc[3, 0]
        print("ToDate",to_date_str)
        from_date_cleaned = from_date_str.replace("From :", "").strip()
        to_date_cleaned = to_date_str.replace("To :", "").strip()

       
        from_date = datetime.strptime(from_date_cleaned, '%d %b %Y')
        to_date = datetime.strptime(to_date_cleaned, '%d %b %Y')
        date_difference = to_date - from_date

     
        days_difference = date_difference.days
        print("FromDate:", from_date)
        print("ToDate:", to_date)
        while row_index < len(df)-9:
           
            try:
                employee_code = df.iloc[row_index, employee_code_col]
            except IndexError:
                print(f"IndexError at row {row_index}: check if data is available")
                break 
           
            for i in range(days_difference+1):
                current_day = from_date + timedelta(days=i)
                if current_day:  
                    attendance_record = {
                        'EmployeeCode': employee_code,
                        'Date': current_day.strftime('%Y-%m-%d'), 
                        'IN': df.iloc[row_index -1, 6+i] or None,
                        'OUT': df.iloc[row_index , 6+i] or None,
                        'DutyHours': df.iloc[row_index +1, 6+i] or None,
                        'STATUS': df.iloc[row_index+2 , 6+i] or None,
                    }
                    attendance_data.append(attendance_record)
                    EmployeeCode =  attendance_record['EmployeeCode']
                    AttendaceDate  = attendance_record['Date']
                    In_Time = attendance_record['IN']
                    Out_Time = attendance_record['OUT']
                    Duty_Hour = attendance_record['DutyHours']
                    Status = attendance_record['STATUS']

                    if EmployeeCode and AttendaceDate:
                           PreviousRecord  = Attendance_Data.objects.filter(OrganizationID=OrganizationID,IsDelete=False,EmployeeCode=EmployeeCode,Date=AttendaceDate)
                           if PreviousRecord: 
                                for Pr in PreviousRecord:
                                    Pr.IsDelete = True
                                    Pr.ModifyBy = UserID
                                    Pr.save()
                           NewRecord =  Attendance_Data.objects.create(OrganizationID=OrganizationID,IsDelete=False,EmployeeCode=EmployeeCode,Date=AttendaceDate,CreatedBy=UserID,In_Time= In_Time,Out_Time= Out_Time,Duty_Hour=Duty_Hour,Status=Status,IsUpload=True)
                           if NewRecord:
                               print(NewRecord.EmployeeCode)

                    

           
            row_index += 5  

        

    return redirect("AttendaceMonthlyReport")

# Attendance  
@transaction.atomic     
def AttendanceProcess(request):
    if 'OrganizationID' not in request.session:
        return redirect(MasterAttribute.Host)
    else:
        print("Show Page Session")

    OrganizationID = request.session["OrganizationID"]
    UserID = str(request.session["UserID"])
    # emp_list =  EmployeeMaster.objects.filter(OrganizationID=OrganizationID,IsDelete=False).order_by('EmployeeCode')
    emp_list = EmployeeDataSelect(OrganizationID)
   
    try:
        org_obj = Organization_Details.objects.get(OID=OrganizationID, IsDelete=False)
        S_EndDate = org_obj.EndDate
    except Organization_Details.DoesNotExist:
        S_EndDate = 26

    current_date = datetime.now().date()
    #current_date =  datetime.strptime("2024-04-26", "%Y-%m-%d").date()
    current_date= datetime(current_date.year, current_date.month, 25).date()

    _, last_day_of_month = calendar.monthrange(current_date.year, current_date.month)

    if S_EndDate is not None:
        safe_day = min(S_EndDate, last_day_of_month)
        start_date = current_date.replace(day=safe_day)
    else:
        safe_day = min(26, last_day_of_month)
        start_date = current_date.replace(day=safe_day)

    # _, last_day_of_month = calendar.monthrange(current_date.year, current_date.month)
    end_date = current_date.replace(day=last_day_of_month)
    with transaction.atomic():
            if request.method == "POST":
                all_emp_codes = request.POST.getlist('all_emp_codes[]')
                all_emp_ids  = request.POST.getlist('all_emp_id[]')
                start_date = datetime.strptime(request.POST.get('Start_Date'), "%Y-%m-%d").date()
                To_Date = datetime.strptime(request.POST.get('To_Date'), "%Y-%m-%d").date()
                end_date=To_Date
                current_date = start_date
                while current_date <= end_date:
                    # for empcode in all_emp_codes:
                    #     if empcode is not None and empcode != '':
                    #         Emp_code = empcode
                    #         # Check if the current date is between start_date and end_date
                    #         if start_date <= current_date <= end_date:
                    #             Att_obj = SalaryAttendance.objects.update_or_create(
                    #                 EmployeeCode=Emp_code,
                    #                 Date=current_date,
                    #                 In_Time='00:00',
                    #                 Out_Time='00:00',
                    #                 Duty_Hour='00:00',
                    #                 OrganizationID=OrganizationID,
                    #                 CreatedBy=UserID,
                    #                 Status='Present',
                    #                 IsUpload=False,
                    #                 IsDelete=False
                    #             )
                    
                    for emp_code, emp_id in zip(all_emp_codes, all_emp_ids):
                        if emp_code:
                            Att_obj = SalaryAttendance.objects.update_or_create(
                                EmployeeCode=emp_code,
                                EmployeeID=emp_id,
                                Date=current_date,
                                In_Time='00:00',
                                Out_Time='00:00',
                                Duty_Hour='00:00',
                                OrganizationID=OrganizationID,
                                CreatedBy=UserID,
                                Status='Present',
                                IsUpload=False,
                                IsDelete=False
                            )

                    current_date += timedelta(days=1)

                messages.success(request,"Attendace Record Created Successfully !")
                return redirect('Daily_Attendace')
        
    context = {
        'Emp_list': emp_list, 
        'Start_Date': start_date,
        'To_Date': end_date,
    }
    return render(request, "EMP_PAY/Attendance/AttendanceProcess.html", context)









def Upload_Attendance_List(request):
    if 'OrganizationID' not in request.session:
        return redirect(MasterAttribute.Host)
    else:
        print("Show Page Session")
     
    OrganizationID = request.session["OrganizationID"]
    UserID = str(request.session["UserID"])
    EmployeeCode = request.session["EmployeeCode"] 
    
    current_date = datetime.now()
    
    year = int(request.GET.get('year', current_date.year))
    month_no = int(request.GET.get('month_no', current_date.month))
    
    previous_month = current_date.replace(year=year, month=month_no, day=1) - timedelta(days=1)
    
   
    next_month = min(
        current_date.replace(year=year, month=month_no, day=28) + timedelta(days=4),
        datetime.now().replace(day=1)
    )
    
    if next_month.month == month_no:
        next_month = next_month.replace(day=1)
    
   
    Data_File = Raw_Attendance_Data_File.objects.filter(
          OrganizationID = OrganizationID,
                       
       
        Attendance_Date__year=year,
        Attendance_Date__month=month_no,
         Attendance_Date__lte=current_date,
        IsDelete=False
    ).order_by('Attendance_Date')
    
   
          
    context = {
        'Data_File': Data_File,
        'current_month': datetime(year, month_no, 1),
        'previous_month': previous_month,
        'next_month': next_month
    }
    
    return render(request, "EMP_PAY/Attendance/Upload_Attendance_List.html", context)





def Payroll_List(request):
    if 'OrganizationID' not in request.session:
        return redirect(MasterAttribute.Host)
    else:
        print("Show Page Session")
     
    OrganizationID = request.session["OrganizationID"]
    UserID = str(request.session["UserID"])
    EmployeeCode = request.session["EmployeeCode"]
    
    current_year = datetime.now().year
    year = int(request.GET.get('current_year', current_year))
    
    action = request.GET.get('action')
    if action == 'prev_year':
        year -= 1
    elif action == 'next_year':
        year += 1
    
    months_data = SalarySlip.objects.filter(
                    OrganizationID=OrganizationID, 
                    IsDelete=False, 
                    EmployeeCode=EmployeeCode,
                    generated = True,
                    HrVerify = True,
                    FcVerify =  True,
                    year=year
                ).order_by('month')

    context = {
        'months_data': months_data,
        'current_year': year,
    }
    
    return render(request, "EMP_PAY/Salary/Payroll_List.html", context)

 

from collections import defaultdict
    
from .models import SalarySlip
from collections import defaultdict
from datetime import datetime, timedelta
from django.shortcuts import render, redirect


def Employees_Payroll_List(request):
    if 'OrganizationID' not in request.session:
        return redirect(MasterAttribute.Host)
    else:
        print("Show Page Session")

    OrganizationID = request.session["OrganizationID"]
    UserID = str(request.session["UserID"])
    UserType = request.session["UserType"]
    current_date = datetime.now()

    year = int(request.GET.get('year', current_date.year))
    month_no = int(request.GET.get('month_no', current_date.month))

    emp_list = EmployeeDataSelectForSalary(OrganizationID,month_no,year)

    # for emp in emp_list:
    #     if emp['EmployeeCode'] == '567':
    #         # print(f"The Employee is here:: {emp['EmployeeCode']}")
    #         pass
    #     else:
    #         # print("The data is not found")
    #         pass

    previous_month = current_date.replace(year=year, month=month_no, day=1) - timedelta(days=1)
    next_month = min(
        current_date.replace(year=year, month=month_no, day=28) + timedelta(days=4),
        datetime.now().replace(day=1)
    )
    if next_month.month == month_no:
        next_month = next_month.replace(day=1)

    salary_slips = SalarySlip.objects.filter(
        EmployeeCode__in=[emp['EmployeeCode'] for emp in emp_list],
        month=month_no,
        year=year,
        OrganizationID=OrganizationID,
        IsDelete=False
    )

    attendancelock = AttendanceLock.objects.filter(
        EmployeeCode__in=[emp['EmployeeCode'] for emp in emp_list],
        month=month_no,
        year=year,
        OrganizationID=OrganizationID,
        IsDelete=False
    ).values('EmployeeCode', 'IsLock', 'id')

    generated_dict = defaultdict(lambda: {
        'generated': False,
        'id': 0,
        'approval_html': '',
        'status': 'Pending'
    })
    lock_dict = defaultdict(lambda: {'IsLock': False, 'id': 0})

    for slip in salary_slips:
     
        generated_dict[slip.EmployeeCode] = {
            'generated': slip.generated,
            'id': slip.id,
            'approval_html': slip.get_approval_status_html()  ,  
            'status': "Approved" if slip.HrVerify and slip.FcVerify else "Pending"  
        }

    for lock in attendancelock:
        lock_dict[lock['EmployeeCode']] = {'IsLock': lock['IsLock'], 'id': lock['id']}

    for emp in emp_list:
        slip_info = generated_dict[emp['EmployeeCode']]
        lock_info = lock_dict[emp['EmployeeCode']]

        emp['generated'] = slip_info['generated']
        emp['id'] = slip_info['id']

        emp['IsLock'] = lock_info['IsLock']
        emp['lockid'] = lock_info['id']
        emp['get_approval_status_html'] = slip_info['approval_html']  or '<p style="color: #f8ac59;">Attendacne Lock is  Pending</p>'
        emp['status'] = slip_info['status']  

    context = {
        'emps': emp_list,
        'current_month': datetime(year, month_no, 1),
        'previous_month': previous_month,
        'next_month': next_month,
        'emp_list': emp_list,
        'month_no': month_no,
        'year': year
    }

    return render(request, "EMP_PAY/Salary/Employees_Payroll_List.html", context)


# These functions are moved to leave
def View_Attendance(request):
    if 'OrganizationID' not in request.session:
        return redirect(MasterAttribute.Host)
    else:
        print("Show Page Session")
     
    OrganizationID = request.session["OrganizationID"]
    UserID = str(request.session["UserID"])
    EmployeeCode = request.session["EmployeeCode"] 
    
    current_date = datetime.now()
    
    year = int(request.GET.get('year', current_date.year))
    month_no = int(request.GET.get('month_no', current_date.month))
    
    previous_month = current_date.replace(year=year, month=month_no, day=1) - timedelta(days=1)
    
   
    next_month = min(
        current_date.replace(year=year, month=month_no, day=28) + timedelta(days=4),
        datetime.now().replace(day=1)
    )
    
    if next_month.month == month_no:
        next_month = next_month.replace(day=1)
    
   
    attendance_data = SalaryAttendance.objects.filter(
        OrganizationID=OrganizationID, 
        EmployeeCode=EmployeeCode,
        Date__year=year,
        Date__month=month_no,
         Date__lte=current_date,
        IsDelete=False
    ).order_by('Date')
    
    for record in attendance_data:

        update_request_exists = Update_Attendance_Request_V1.objects.filter(
            SalaryAttendance=record,
            OrganizationID=OrganizationID, 
            IsDelete=False 
        ).exists()
        record.update_request_exists = update_request_exists
          
    context = {
        'attendance_data': attendance_data,
        'current_month': datetime(year, month_no, 1),
        'previous_month': previous_month,
        'next_month': next_month
    }
    
    return render(request, "EMP_PAY/Attendance/View_Attendance.html", context)



def Daily_Attendace(request):
    if 'OrganizationID' not in request.session:
        return redirect(MasterAttribute.Host)
    else:
        print("Show Page Session")
     
    OrganizationID = request.session["OrganizationID"]
    UserID = str(request.session["UserID"])
    
    emp_list = EmployeeDataSelect(OrganizationID)
    
    
    Q_Date  =  request.GET.get('Q_Date')

    attendance_data = SalaryAttendance.objects.filter(
        OrganizationID=OrganizationID, 
        Date= Q_Date,
        IsDelete=False
    ).order_by('Date','EmployeeCode')
    
    # for record in attendance_data:
    #     update_request_exists = Update_Attendance_Request.objects.filter(Attendance_Data=record,OrganizationID=OrganizationID, 
    #                                         IsDelete=False ).exists()
    #     record.update_request_exists = update_request_exists

    for record in attendance_data:
        for emp in emp_list:
            if emp['EmployeeCode'] == record.EmployeeCode:
                
                record.EmployeeName = emp['EmpName']
                break
        
        # update_request_exists = Update_Attendance_Request.objects.filter(
        #     SalaryAttendance=record,
        #     OrganizationID=OrganizationID,
        #     IsDelete=False
        # ).exists()
        
        update_request_exists = Update_Attendance_Request_V1.objects.filter(
            SalaryAttendance=record,
            OrganizationID=OrganizationID,
            IsDelete=False
        ).exists()
        record.update_request_exists = update_request_exists

          
    context = {
        'attendance_data': attendance_data,
        'Q_Date':Q_Date
    }
    
    return render(request, "EMP_PAY/Attendance/Daily_Attendace.html", context)

# Ends Here------------

def get_month_range(year, month_no):
    start_date = datetime(year, month_no, 1)
    next_month = start_date.replace(day=28) + timedelta(days=4)
    end_date = next_month - timedelta(days=next_month.day)
    return start_date, end_date

def get_previous_month_start_date(year, month_no):
    if month_no == 1:
        year -= 1
        month_no = 12
    else:
        month_no -= 1
    return datetime(year, month_no, 26)

def get_attendance_counts(attendance, leavelist):
    week_off_count = 0
    present_count = 0
    absent_count = 0
    leave_counts = {leave: 0 for leave in leavelist}

    for lv in leavelist:
        print(f"the lv is here :: {lv}")

    for att in attendance:
        status = att.Status
        if status.lower()  == 'week off':
            week_off_count += 1
        elif status.lower() == 'present':
            present_count += 1
        elif status.lower() == 'absent' or status is None:
            absent_count += 1
        elif status in leavelist:
            leave_counts[status] += 1

    leave_counts_Demo = sum(leave_counts.values())
    # print("--------------- get_attendance_counts --------------- ")
    # print(f"leave_counts == {leave_counts}")
    # print(f"week_off_counts == {week_off_count}")
    # print(f"present_count == {present_count}")
    # print(f"leave_counts_Demo == {leave_counts_Demo}")
    # print("--------------- / get_attendance_counts --------------- ")


    total_working_days = week_off_count + present_count + sum(leave_counts.values())
    return week_off_count, present_count, absent_count, leave_counts, total_working_days

def calculate_paid_days(total_working_days, total_days, total_no_days_in_month):
    if total_working_days>total_no_days_in_month:
        adjustment = len(total_days) - total_no_days_in_month
        return total_working_days - adjustment
    return total_working_days

def days_in_selected_month(month_no, year):
    if month_no == 12:
        next_month = 1
        next_month_year = year + 1
    else:
        next_month = month_no + 1
        next_month_year = year
    start_date = datetime(year, month_no, 1)
    end_date = datetime(next_month_year, next_month, 1)
    return (end_date - start_date).days
import calendar


def AttendanceLockView(request):
    if 'OrganizationID' not in request.session:
        return redirect(MasterAttribute.Host)
    
    OrganizationID = request.session["OrganizationID"]
    UserID = str(request.session["UserID"])
    EmployeeCode = request.GET.get('emp')
    emp_Details = EmployeeDataSelect(OrganizationID, EmployeeCode)

    year = int(request.GET.get('year'))
    month_no = int(request.GET.get('month_no'))
    month_name = calendar.month_name[month_no]

    # Correctly calculate the start and end dates
    start_date = get_previous_month_start_date(year, month_no)
    end_date = start_date.replace(day=28) + timedelta(days=4)
    end_date = end_date.replace(day=25)
    # od = Organization_Details.objects.filter(OID=OrganizationID, IsDelete=False).first()

    # print("the od values is here::", od)
    # print("the od.EndDate values is here::", od.EndDate)
    
    # if od:
    #     if od.EndDate is not None and od.EndDate==1:
    #         start_date = datetime(year, month_no, 1)

    #         # Get last day of the month using calendar
    #         last_day = calendar.monthrange(year, month_no)[1]

    #         # End date is the last day of the month
    #         end_date = datetime(year, month_no, last_day)

    od = Organization_Details.objects.filter(OID=OrganizationID, IsDelete=False).first()

    if od:
        if od.EndDate == 1:  
            # Cycle = 1st → Last Day of Month
            start_date = datetime(year, month_no, 1)
            last_day = calendar.monthrange(year, month_no)[1]
            end_date = datetime(year, month_no, last_day)

        elif od.EndDate == 31:
            # Cycle = 1st → 31st (normal calendar month)
            start_date = datetime(year, month_no, 1)
            last_day = calendar.monthrange(year, month_no)[1]  # handles Feb (28/29), April (30), etc.
            end_date = datetime(year, month_no, last_day)

        else:
            # Default cycle: 26th prev month → EndDate current month
            start_date = get_previous_month_start_date(year, month_no)
            end_date = datetime(year, month_no, od.EndDate)

    total_days = [(start_date + timedelta(days=i)).strftime('%#d') for i in range((end_date - start_date).days + 1)]
    date_list = [start_date + timedelta(days=x) for x in range((end_date - start_date).days + 1)]

            
    # Generate total_days list
    # total_days = [(start_date + timedelta(days=i)).strftime('%#d') for i in range((end_date - start_date).days + 1)]

    #attendance = Attendance_Data.objects.filter(EmployeeCode=EmployeeCode, OrganizationID=OrganizationID, IsDelete=False, Date__range=(start_date, end_date)).order_by('Date')
    # date_list = [start_date + timedelta(days=x) for x in range((end_date - start_date).days + 1)]
    
    # for date in date_list:
    #     r = Attendance_Data.objects.filter(EmployeeCode=EmployeeCode, OrganizationID=OrganizationID, IsDelete=False, Date=date)
    #     if r.exists():
    #         # print()
    #         pass
    #     else:
    #         Attendance_Data.objects.create(Date=date,Status="Absent",OrganizationID=OrganizationID,EmployeeCode=EmployeeCode)
    #         #complete_attendance.append(attendance_dict[date])  # Existing record
    
    attendance = Attendance_Data.objects.filter(
    # attendance = SalaryAttendance.objects.filter(
        EmployeeCode=EmployeeCode, 
        OrganizationID=OrganizationID, 
        IsDelete=False, 
        Date__range=(start_date, end_date)
    ).order_by('Date')

    leave_type = Leave_Type_Master.objects.filter(IsDelete=False, Is_Active=True)
    leavelist = [leave.Type for leave in leave_type]

    # print("attendence data ::", attendance)

    lockid = request.GET.get('lockid')
    IsLock = request.GET.get('IsLock')

    if lockid and lockid != '0':
        lock_obj = get_object_or_404(AttendanceLock, id=lockid, IsDelete=False)
    # print(len(total_days))

    if request.method == "POST":
        btn = request.POST['btn']
        IsLock = btn == "Lock"

        week_off_count, present_count, absent_count, leave_counts, total_working_days = get_attendance_counts(attendance, leavelist)
        # total_no_days_in_month = days_in_selected_month(month_no, year)
        total_no_days_in_month = len(date_list)
        pai_days = calculate_paid_days(total_working_days, total_days, total_no_days_in_month)
        total_days_count = len(total_days)
        # print("-----------------------------------------------")
        # print(f"total_working_days = {total_working_days}")
        # print(f"total_days({total_days_count}) :=: {total_days}")
        # print(f"total_no_days_in_month = {total_no_days_in_month}")
        # print(f"pai_days = {pai_days}")
        # print("-----------------------------------------------")

        # print(f"start_date = {start_date}, end_date = {end_date}, total_days = {len(total_days)}")

        if lockid and lockid != '0':
            lock_obj.PaiDays = pai_days
            lock_obj.ModifyBy = UserID
            lock_obj.IsLock = IsLock
            lock_obj.save()
        else:
            AttendanceLock.objects.create(
                OrganizationID=OrganizationID,
                CreatedBy=UserID,
                EmployeeCode=EmployeeCode,
                month=month_no,
                month_name=month_name,
                year=year,
                PaiDays=pai_days,
                total_no_Days_in_month=total_no_days_in_month,
                IsLock=IsLock
            )

        messages.success(request, "Saved Successfully" if btn == "Save" else "Locked Successfully")
        return redirect(reverse('Employees_Payroll_List') + f'?year={year}&month_no={month_no}')

    unique_attendance = []

    for date, records in groupby(attendance, key=attrgetter("Date")):
        # Convert the groupby iterator to a list for safe handling
        records_list = list(records)
        
        # Prioritize "Present" if it exists
        present_record = next((r for r in records_list if r.Status == "Present"), None)
        
        # Add either the "Present" record or the first available record if it exists
        if present_record:
            unique_attendance.append(present_record)
        elif records_list:
            unique_attendance.append(records_list[0])  

    context = {
        'attendance_data': unique_attendance,
        'leave_type': leave_type,
        'emp_Details': emp_Details,
        'month': month_name,
        'year': year,
        'lockid': lockid,
        'IsLock': IsLock
    }
    return render(request, "EMP_PAY/Attendance/AttendanceLockView.html", context)



def UnlockAttendanceLock(request):
    if 'OrganizationID' not in request.session:
        return redirect(MasterAttribute.Host)
    
    OrganizationID = request.session["OrganizationID"]
    UserID = str(request.session["UserID"])
    EmployeeCode = request.GET.get('emp')
    emp_Details = EmployeeDataSelect(OrganizationID, EmployeeCode)

    year = int(request.GET.get('year'))
    month_no = int(request.GET.get('month_no'))
    month_name = calendar.month_name[month_no]

    od = Organization_Details.objects.filter(OID=OrganizationID, IsDelete=False).first()

    lockid = request.GET.get('lockid')
    IsLock = request.GET.get('IsLock')

    if lockid and lockid != '0':
        lock_obj = get_object_or_404(AttendanceLock, id=lockid, IsDelete=False)
    
        if lockid and lockid != '0':
            lock_obj.ModifyBy = UserID
            lock_obj.IsLock = False
            lock_obj.save()
        
        
        return redirect(reverse('Employees_Payroll_List') + f'?year={year}&month_no={month_no}')


 

from datetime import datetime, timedelta
from calendar import monthrange


from Leave_Management_System.models import Emp_Leave_Balance_Master,EmpMonthLevelCreditMaster


def UpdateStatus(request):
    if 'OrganizationID' not in request.session:
        return redirect(MasterAttribute.Host)
    else:
        print("Show Page Session")
     
    OrganizationID = request.session["OrganizationID"]
    UserID = str(request.session["UserID"])
    UserFullName = request.session["FullName"]

    if request.method == "GET":
        Attid = request.GET.get('Attid')
        Status = request.GET.get('Status')
        Remarks = request.GET.get('Remarks') or ''
        LeaveID = request.GET.get('leaveTypeId') or 0

        # print("the leave id is here::", LeaveID)
        # print("the Remarks is here::", Remarks)
          
        # att_obj = get_object_or_404(Attendance_Data,id=Attid,IsDelete=False,OrganizationID=OrganizationID)
        att_obj = get_object_or_404(SalaryAttendance,id=Attid,IsDelete=False,HotelID=OrganizationID)

        non_present_statuses = ["Absent", "LWP"]

        if Status == "Half Day Present":
            att_obj.IsPresent = True
            att_obj.PresentValue = 0.5
        elif Status not in non_present_statuses:
            att_obj.IsPresent = True
            att_obj.PresentValue = 1
        else:
            att_obj.IsPresent = False
            att_obj.PresentValue = 0

        att_obj.Status = Status
        att_obj.IsAttendanceModified = True
        att_obj.Remarks = Remarks
        att_obj.ModifyBy = UserID
        att_obj.ModifyDateTime = timezone.now()
        att_obj.ActionByName = UserFullName if UserFullName else ''
        att_obj.save()

        response_data = {
            'message': f'Status {Status} updated successfully'
        }
        return JsonResponse(response_data, status=200)
        
        
        

def get_employee_attendance_data():
    try:
        with connection.cursor() as cursor:
            # Execute the stored procedure
            cursor.execute("{CALL getEmpAttendanceData}")

            # Fetch all rows (assuming a result set is returned)
            results = cursor.fetchall()

            # Manually specify the column names if `description` is unavailable
            columns = ["CID","ID","NAME","DEPARTMENT", "DESIGNATION","EmployeeCode","Date", "In_Time", "Out_Time", "S_In_Time", "S_Out_Time", "Duty_Hour", "Status"]

            # Map the results to a list of dictionaries
            data = [dict(zip(columns, row)) for row in results]

            return data
    except Exception as e:
        print(f"Error fetching data: {e}")
        return []
import json
from collections import defaultdict
import datetime
from datetime import date
def process_attendance_data(raw_data):
    # Example function to process the data and convert date objects to strings
    for record in raw_data:
        if isinstance(record.get('Date'), date):
            record['Date'] = record['Date'].strftime('%Y-%m-%d')  # Convert date to string
        # If you have any other date fields, you can handle them similarly
        if isinstance(record.get('In_Time'), date):
            record['In_Time'] = record['In_Time'].strftime('%Y-%m-%d')
        if isinstance(record.get('Out_Time'), date):
            record['Out_Time'] = record['Out_Time'].strftime('%Y-%m-%d')

    return raw_data

from datetime import datetime, timedelta
from django.shortcuts import render

def attendance_report(request):
    report_data = get_employee_attendance_data()  # This should return a list of records
    
    start_date = datetime.strptime(request.GET.get('start_date', '2024-10-26'), '%Y-%m-%d')
    end_date = datetime.strptime(request.GET.get('end_date', '2024-11-25'), '%Y-%m-%d')

    # Group the attendance data by EmployeeCode
    # grouped_data = {}
    # for record in report_data:
    #     emp_code = record['EmployeeCode']
    #     if emp_code not in grouped_data:
    #         grouped_data[emp_code] = []
    #     grouped_data[emp_code].append(record)

    grouped_data = {}  # Initialize an empty dictionary to hold grouped records
    for record in report_data:
        emp_code = record['EmployeeCode']  # Get the EmployeeCode for each record
        
        # If this EmployeeCode does not exist in grouped_data, create a new entry with an empty list
        if emp_code not in grouped_data:
            grouped_data[emp_code] = []
        
        # Append the current record to the list of that EmployeeCode
        grouped_data[emp_code].append(record)
    

    # Generate the list of all unique dates in the report data
    all_dates = sorted(set(record['Date'] for record in report_data))

    # Generate the date range from start_date to end_date
    date_range = []
    current_date = start_date
    while current_date <= end_date:
        date_range.append(current_date.strftime('%Y-%m-%d'))
        current_date += timedelta(days=1)

    # Create the table header row with both date and weekday
    header_row = ['NAME',  'EmployeeCode']
    for date_str in date_range:
        date_obj = datetime.strptime(date_str, '%Y-%m-%d')
        day_of_month = date_obj.day
        day_of_week = date_obj.strftime('%A')  # Full weekday name (e.g., "Monday")
        header_row.append(f'{day_of_month} <br/> {day_of_week}')

    # Create the table body for each employee
    table_rows = []
    for emp_code, employee_data in grouped_data.items():
        employee = employee_data[0]  # Assume the first record has all the employee details
        row = [
            employee['NAME'],
            # employee['DEPARTMENT'],
            # employee['DESIGNATION'],
            emp_code+"<br/> In <br/> Out </br>  In <br/> Out <br> Hours </br> Status"
        ]

        # For each date, check if the employee has attendance data
        for date_str in date_range:
            # Find the record for this employee and date
            date_obj = datetime.strptime(date_str, "%Y-%m-%d").date()
            record = next((r for r in employee_data if r['Date'] == date_obj), None)
            
            if record:
                # Format times (remove unnecessary microseconds)
                in_time = record['In_Time'][:-7] if record['In_Time'] else ''
                out_time = record['Out_Time'][:-7] if record['Out_Time'] else ''
                s_in_time = record['S_In_Time'][:-7] if record['S_In_Time'] else ''
                s_out_time = record['S_Out_Time'][:-7] if record['S_Out_Time'] else ''
                duty_hour = record['Duty_Hour'][:-7] if record['Duty_Hour'] else ''
                status = record['Status']
                
                row.append(f'<br/> {in_time} <br/> {out_time} <br/>  {s_in_time}  <br/> {s_out_time} <br/> {duty_hour} <br/><b>{status}</b>')
            else:
                row.append('')  # Empty cell if no data for that date

        table_rows.append(row)

    # Pass the grouped data, date range, and header row to the template
    return render(request, 'EMP_PAY/Attendance/attendance_report.html', {
        'date_range': date_range,
        'header_row': header_row,
        'table_rows': table_rows,
        'start_date': start_date,
        'end_date': end_date
    })


def attendance_report_Old(request):
    raw_data = get_employee_attendance_data()  # This fetches raw data
    raw_data = process_attendance_data(raw_data)  # Convert any date fields to strings

    # Now serialize the data into JSON
    raw_data_json = json.dumps(raw_data)
    start_date = '2024-10-26'
    end_date  = '2024-11-25'
    # Pass the JSON data to the template
    return render(request, 'EMP_PAY/Attendance/attendance_report.html', {'report_data_json': raw_data_json,'start_date':start_date,'end_date':end_date})


from app.models import DepartmentMaster
from django.db import connection
def AttendaceMonthlyReport(request):
    if 'OrganizationID' not in request.session:
        return redirect(MasterAttribute.Host)
    else:
        print("Show Page Session")
     
    OrganizationID = request.session["OrganizationID"]
    UserID = str(request.session["UserID"])
    
    Departments =  DepartmentMaster.objects.filter(IsDelete =False)

    I = request.GET.get('I',OrganizationID)
    GetDepartments = request.GET.get('Departments', 'All')
    # print('GetDepartments', GetDepartments)

    year = request.GET.get('year')
    if year:
        year = int(year)
    else:
        year = datetime.now().year
        
    month_no =  request.GET.get('month_no')
    if month_no:
        month_no = int(month_no)
    else:
        month_no = datetime.now().month  

    emp_list =  EmployeeDataSelectForSalary(OrganizationID,month_no,year)
    EmployeeCode = request.GET.get('EmployeeCode','All')  # Get the EmployeeCode from the form

    # print("EmployeeCode = ",EmployeeCode)
    od =Organization_Details.objects.filter(OID=OrganizationID,IsDelete=False)

    cyS=None
    try:
        if od.exists():
            cyS=od.first().EndDate
    except:
        cyS=None
   
    month_name = calendar.month_name[int(month_no)]    
    StartDate = datetime(year, month_no, 1)


    if cyS is not None and cyS == 1:
        # Cycle = 1st → last day of month
        StartDate = datetime(year, month_no, 1)
        _, last_day = calendar.monthrange(year, month_no)
        EndDate = datetime(year, month_no, last_day)

    elif cyS is not None and cyS == 31:
        # Cycle = full calendar month (1 → 31 / or last day if <31)
        StartDate = datetime(year, month_no, 1)
        _, last_day = calendar.monthrange(year, month_no)
        EndDate = datetime(year, month_no, last_day)

    else:
        # Default cycle: 26th prev month → 25th current month
        if month_no == 1:
            year -= 1
            month_no = 12
        else:
            month_no -= 1
        StartDate = datetime(year, month_no, 26)
        EndDate = datetime(year, month_no + 1, 25) if month_no < 12 else datetime(year + 1, 1, 25)


        next_month = StartDate.replace(day=28) + timedelta(days=4)
        EndDate = next_month.replace(day=25)
        if month_no==12:
            month_no=1
        else:
            month_no +=  1

    # print("----------------------------------------------------------------------::")        
    # print("The End Date is here::", EndDate)        
    # print("The Start Date is here::", StartDate)        
    days = [(StartDate + timedelta(days=i)).strftime('%#d') for i in range((EndDate - StartDate).days + 1)]

    # print("the start date is here::", StartDate)
    # print("the Start Date is here::", EndDate)
    # print("the EmployeeCode is here::", EmployeeCode)
    # print("the OrganizationID is here::", OrganizationID)
    
    with connection.cursor() as cursor:
        # cursor.execute("EXEC GetAttendancePivot_New %s, %s, %s, %s", [OrganizationID, StartDate, EndDate, EmployeeCode])
        cursor.execute("EXEC GetAttendancePivot_New_For_Payroll %s, %s, %s, %s", [OrganizationID, StartDate, EndDate, EmployeeCode])
        rows = cursor.fetchall()
        columns = [col[0] for col in cursor.description]
    
    # print("----------------------------------------------------------------------::")    
    # print("OrganizationID::",OrganizationID)    
    # print("StartDate::",StartDate)    
    # print("EndDate::",EndDate)    
    # print("EmployeeCode::",EmployeeCode)    
    # print("----------------------------------------------------------------------::")        


    rowslist = [dict(zip(columns, row)) for row in rows]
    
   
    leavetype = Leave_Type_Master.objects.filter(IsDelete=False,Is_Active=True)
    leavelist = [leave.Type for leave in leavetype]
   
    for row in rowslist:
            TotalDays = len(days)
            Presentco=0
            WeekOffCount = 0
            PresentCount = 0
            AbsentCount = 0
            Total_AR=0
            l_Count = 0
            leave_counts = {leave: 0 for leave in leavelist}
            
            for day in days:
                status = row[day]
                # print("staus is here::", status)
                actual_status=None
                if status:
                    # Extract the actual status from the last part of the string
                    actual_status = status.split("^")[-1].strip()
                    # print("actual_status is here::", actual_status)
                if actual_status is not None:
                    if actual_status.lower() == 'week off' or actual_status.lower()=='w': 
                        WeekOffCount += 1
                    elif actual_status.lower() == 'p' or actual_status.lower() == 'present':
                        PresentCount += 1
                    elif actual_status.lower() == 'co' or actual_status.lower() == 'comp-off':
                        Presentco += 1
                        Aactual_status="Comp-off"
                        leave_counts[Aactual_status] =Presentco
                    elif actual_status.lower() == 'ar':
                        Total_AR += 1 
                    elif actual_status.lower() == 'a' or actual_status == None :
                        AbsentCount += 1
                       
                    if actual_status in leavelist:
                        if actual_status.lower() != 'ar':
                            if actual_status.lower() == 'co' or actual_status.lower() == 'comp-off':
                                actual_status="Comp-off"
                                # leave_counts[actual_status] =Presentco
                            else:
                                l_Count=l_Count+1
                                leave_counts[actual_status] += 1
                        if actual_status.lower() == 'ar':
                            leave_counts[actual_status] += 1

            # for leave in leavelist:
              
            #     l_Count =  l_Count + leave_counts[leave] 
            TotalWorkingDays = WeekOffCount + PresentCount + l_Count+Total_AR+Presentco
            row['TotalWorkingDays'] =  TotalWorkingDays
            row['Present'] = PresentCount
            row['Absent'] = AbsentCount
            row['WeekOff'] = WeekOffCount
            row["Comp-off"]=Presentco
            total_no_Days_in_month = days_in_selected_month(int(month_no), int(year))
            minus = TotalDays - total_no_Days_in_month 
            if row["iscalm"]=="1":
                row['TotalPaidDays'] = TotalWorkingDays - ( minus )
            else:
                row['TotalPaidDays'] = TotalWorkingDays
            

            row['leave_counts'] = leave_counts 
            row['l_Count'] = l_Count 
            # row['TotalLeave'] = l_Count
            # print("leave counts::", l_Count)

        
        
    today = datetime.today()
    CYear = today.year
    CMonth = today.month      

    context = {
        'CYear':range(CYear,2020,-1),'CMonth':CMonth,
        'month_no':month_no,
        'year':year,
        'month_name':month_name,
        'rowslist':rowslist,
        'days':days,
        'leavelist':leavelist,
        'emp_list':emp_list,
        'EmployeeCode':EmployeeCode,
        'Departments':Departments,
        'StartDate':StartDate,
        'EndDate':EndDate,
        'EmployeeCode':EmployeeCode,
        'OrganizationID':OrganizationID,
        'GetDepartments':GetDepartments,
    }
     
    return render(request, "EMP_PAY/Attendance/AttendaceMonthlyReport.html", context)





def UpdateRequestlist(request):
    if 'OrganizationID' not in request.session:
        return redirect(MasterAttribute.Host)
    else:
        print("Show Page Session")
     
    OrganizationID = request.session["OrganizationID"]
    UserID = str(request.session["UserID"])
    EmployeeCode = request.session["EmployeeCode"]
    Status = request.GET.get('Status')
    if Status:
        Apporve_status = Status
    if Status is None:
        Apporve_status = 0    

    # att_reqs = Update_Attendance_Request.objects.filter(OrganizationID=OrganizationID,Attendance_Data__EmployeeCode= EmployeeCode,IsDelete=False,Apporve_status= Apporve_status)
    att_reqs = Update_Attendance_Request_V1.objects.filter(OrganizationID=OrganizationID,SalaryAttendance__EmployeeCode= EmployeeCode,IsDelete=False,Apporve_status= Apporve_status)
    
    context = {'att_reqs':att_reqs,'Apporve_status':Apporve_status}
    return render(request, "EMP_PAY/Attendance/UpdateRequestlist.html", context)

 





def UpdateRequestlist_HR(request):
    if 'OrganizationID' not in request.session:
        return redirect(MasterAttribute.Host)
    else:
        print("Show Page Session")
     
    OrganizationID = request.session["OrganizationID"]
    UserID = str(request.session["UserID"])
    # EmployeeCode = request.session["EmployeeCode"]
    Status = request.GET.get('Status')
    
    
    # emp_list =  EmployeeMaster.objects.filter(OrganizationID=OrganizationID, IsDelete=False).order_by('EmployeeCode')
    emp_list =  EmployeeDataSelect(OrganizationID)
    
    
    if Status:
        Apporve_status = Status
    if Status is None:
        Apporve_status = 0    

    # att_reqs = Update_Attendance_Request.objects.filter(OrganizationID=OrganizationID,IsDelete=False,Apporve_status= Apporve_status)
    att_reqs = Update_Attendance_Request_V1.objects.filter(OrganizationID=OrganizationID,IsDelete=False,Apporve_status= Apporve_status)
    # print("att_reqs = ",att_reqs)
  
    employee_data = []
    try:
        for emp in emp_list:
            
            EmployeeCode = emp['EmployeeCode']
            # print
            # Update_Attendance = Update_Attendance_Request.objects.filter(Attendance_Data__EmployeeCode=EmployeeCode,OrganizationID=OrganizationID,IsDelete=False)
            Update_Attendance = Update_Attendance_Request_V1.objects.filter(SalaryAttendance__EmployeeCode=EmployeeCode,OrganizationID=OrganizationID,IsDelete=False)

            if Update_Attendance:
            
                employee_details = {
                    'EmployeeCode': EmployeeCode,
                    'EmpName': emp['EmpName'],
                    'Department': emp['Department'],
                    'Designation': emp['Designation']
                    
                }
                employee_data.append(employee_details)
    except:
        messages.warning(request,'Data Not Found')            
   
    context = {
        'att_reqs':att_reqs,
        'Apporve_status':Apporve_status,
        'employee_data':employee_data
    }
    return render(request, "EMP_PAY/Attendance/UpdateRequestlist_HR.html", context)



@transaction.atomic
def UpdateRequest(request):
    if 'OrganizationID' not in request.session:
        return redirect(MasterAttribute.Host)
    else:
        print("Show Page Session")
     
    OrganizationID = request.session["OrganizationID"]
    UserID = str(request.session["UserID"])
    EmployeeCode = request.session["EmployeeCode"]
    id = request.GET.get('ID')
    obj = Attendance_Data.objects.get(id=id)
    with transaction.atomic():
        if request.method == "POST":
            Reason = request.POST['Reason'] 
            # req =  Update_Attendance_Request.objects.create(
            req =  Update_Attendance_Request_V1.objects.create(
                SalaryAttendance_id =id,
                Reason=Reason,
                CreatedBy= UserID,
                OrganizationID=OrganizationID,
                IsDelete=False
            )
            messages.success(request,"Applied Successfully")
            return redirect('UpdateRequestlist')
    
    context = {'obj':obj} 
    return render(request,"EMP_PAY/Attendance/UpdateRequest.html", context)


from django.shortcuts import render, redirect, get_object_or_404
from django.urls import reverse
from django.contrib import messages
from django.db import transaction

@transaction.atomic
def Update_Attendance(request):
    if 'OrganizationID' not in request.session:
        return redirect(MasterAttribute.Host)
    
    organization_id = request.session["OrganizationID"]
    user_id = str(request.session["UserID"])
    attendance_id = request.GET.get('ID')
    update_request_id = request.GET.get('U_ID')
    
    # Fetch required objects
    obj = get_object_or_404(SalaryAttendance, id=attendance_id)
    # req = get_object_or_404(Update_Attendance_Request, id=update_request_id)
    req = get_object_or_404(Update_Attendance_Request_V1, id=update_request_id)
    pre_status = obj.Status

    if request.method == "POST":
        btn = request.POST.get('btn')
        with transaction.atomic():
            # Update Attendance Data
            obj.In_Time = request.POST.get('in_time') or obj.In_Time
            obj.Out_Time = request.POST.get('out_time') or obj.Out_Time
            obj.S_In_Time = request.POST.get('s_in_time') or obj.S_In_Time
            obj.S_Out_Time = request.POST.get('s_out_time') or obj.S_Out_Time
            new_status = request.POST.get('status')

            if new_status:
                obj.Status = new_status
            else:
                obj.Status = pre_status

            obj.ModifyBy = user_id
            obj.save()

            # Update Request Approval Status
            req.Apporve_status = 1 if btn == "1" else -1
            req.ModifyBy = user_id
            req.save()

            messages.success(request, "Attendance updated successfully.")
            return redirect(reverse('UpdateRequestlist_HR') + f'?Status={btn}')

    context = {
        'obj': obj,
        'req': req,
    }
    return render(request, "EMP_PAY/Attendance/Update_Attendance.html", context)


# Moved To Leave Module
@transaction.atomic
def Update_Attendance_HR(request):
    if 'OrganizationID' not in request.session:
        return redirect(MasterAttribute.Host)
    else:
        print("Show Page Session")
     
    OrganizationID = request.session["OrganizationID"]
    UserID = str(request.session["UserID"])
    id = request.GET.get('ID')
    
    obj = Attendance_Data.objects.get(id=id)
    pre_status = obj.Status
    leave_type  =Leave_Type_Master.objects.filter(Is_Active=True,IsDelete=False)
    
    with transaction.atomic():
        if request.method == "POST":
                Page = request.POST.get('Page')
                emp = request.POST.get('emp')
                year = request.POST.get('year')
                month_no = request.POST.get('month_no')


           
                obj.In_Time = request.POST.get('in') 
                obj.Out_Time =  request.POST.get('Out')
                Q_Date = obj.Date
                
                Status =  request.POST.get('Status')
                if Status == "":
                    obj.Status = pre_status
                else:
                    obj.Status = Status    
                obj.ModifyBy = UserID
                obj.save()
            
                messages.success(request,"Updated Successfully")
                if Page == "VESD":
                    return redirect(reverse('View_Emmployee_Salary_Details') + f'?emp={emp}&year={year}&month_no={month_no}')
                    
                if Page  == "DA":
                    
                    return redirect(reverse('Daily_Attendace') + f'?Q_Date={Q_Date}') 
                
                    

    context ={'obj':obj,'leave_type':leave_type}
    return render(request,"EMP_PAY/Attendance/Update_Attendance_HR.html", context)









def days_in_selected_month(month_no, year):
    num_days = calendar.monthrange(year, month_no)[1]
    return num_days


import pandas as pd


def get_salary_component(sal_details, title):
    return next(
        (item['Monthly'] for item in sal_details if item['Title'].strip().lower() == title.lower()), 
        0
    )


@transaction.atomic
def View_Emmployee_Salary_Details(request):
    if 'OrganizationID' not in request.session:
        return redirect(MasterAttribute.Host)
    else:
        print("Show Page Session")
     
    OrganizationID = request.session["OrganizationID"]

    # print("SESSION ORGANIZATION ID IS HERE::", OrganizationID)

    UserID = str(request.session["UserID"])
    EmployeeCode = request.GET.get('emp')
    org_Details  =  OrganizationMaster.objects.get(OrganizationID= OrganizationID)
    OrgConfig = Organization_Details.objects.filter(OID= OrganizationID,IsDelete=False)
    if OrgConfig.exists():
        OrgConfig=OrgConfig.first()
    else:
        OrgConfig=None
    leave_type  =Leave_Type_Master.objects.filter(Is_Active=True,IsDelete=False)

    hotelapitoken = MasterAttribute.HotelAPIkeyToken
    headers = {
    'hotel-api-token': hotelapitoken  
    }
    api_url = "http://hotelops.in/api/PyAPI/HREmployeeSalary?EmpCode="+str(EmployeeCode)+"&OrganizationID="+str(OrganizationID)

    try:
        response = requests.get(api_url, headers=headers)
        response.raise_for_status()  
        emp_sal_Det = response.json()
        # print("the entire user data:", emp_sal_Det)
    except requests.exceptions.RequestException as e:
         print(f"Error occurred: {e}")
    
    api_ur = "http://hotelops.in/api/PyAPI/HREmployeeDataEmpCodeSelect?EmpCode="+str(EmployeeCode)+"&OID="+str(OrganizationID)

    try:
        response = requests.get(api_ur, headers=headers)
        response.raise_for_status()  
        emp_Details = response.json()
        # print("emp_Details == : : :",emp_Details)
        date_object = datetime.strptime(emp_Details[0]['DateofBirth']  , '%Y-%m-%d')
        formatted_date_DOB = date_object.strftime('%d %b %Y')
        emp_Details[0]['DOB'] = formatted_date_DOB
      
    except requests.exceptions.RequestException as e:
        print(f"Error occurred: {e}")


    if not emp_sal_Det or len(emp_sal_Det) < 12:
        messages.error(request, "Salary details could not be fetched from API.")
        return redirect('Employees_Payroll_List')
    
    # print(emp_Details)
    ID =  request.GET.get('ID')
    sal_obj  = None
    Stored_Arrear  = 0
    if ID != '0' :
        sal_obj = get_object_or_404(SalarySlip,id= ID,IsDelete = False,OrganizationID = OrganizationID)
        # if sal_obj:
        #     Stored_Arrear =  sal_obj.Arrear

        if sal_obj:
            Stored_Arrear = sal_obj.Arrear
        else:
            print("No data found::")


        # print("The Already Stored Arrear::", Stored_Arrear)
        # sal_obj_update = get_object_or_404(SalarySlip,id=ID,IsDelete = False,OrganizationID = OrganizationID)

    year = request.GET.get('year')
    month_no =  request.GET.get('month_no')
    month_name = calendar.month_name[int(month_no)]
   

    fixed_basic_component = get_salary_component(emp_sal_Det, "Basic")
    fixed_HRA_component = get_salary_component(emp_sal_Det, "HRA")
    # Conveyance_Allowance_component = 0
    CCA_component = get_salary_component(emp_sal_Det, "Conveyance Allowance")
    Other_Allowance_component = get_salary_component(emp_sal_Det, "Other Allowance")
    PT_component = get_salary_component(emp_sal_Det, "PT")
    Meals_component = get_salary_component(emp_sal_Det, "Meals")
    Accommodation_component = get_salary_component(emp_sal_Det, "Accommodation")
    TTC_component = get_salary_component(emp_sal_Det, "Total Company Contribution (C)")
    EmployeePF_component = get_salary_component(emp_sal_Det, "Employee PF @12% (Basic)")

    Conveyance_Allowance  = 0

    fixed_basic = fixed_basic_component if fixed_basic_component else 0.0      
    fixed_HRA = fixed_HRA_component if fixed_HRA_component else 0.0      
    CCA = CCA_component if CCA_component else 0.0     
    Other_Allowance = Other_Allowance_component if Other_Allowance_component else 0.0      
    PT = PT_component if PT_component else 0.0      
    Meals = Meals_component if Meals_component else 0.0      
    Accommodation = Accommodation_component if Accommodation_component else 0.0      
    TTC = TTC_component if TTC_component else 0.0     
    EmployeePF = EmployeePF_component if EmployeePF_component else 0.0      

    # print("the basic ammount is here - fixed_basic - :::", fixed_basic)

    fixed_Gross_Per_Annm = next(
        (item['Annum'] for item in emp_sal_Det if item['Title'].startswith("Gross (A)")), 
        0
    )

    # print("the gross of annum", fixed_Gross_Per_Annm)


    BankName =  emp_Details[0]['BankName'] 
    BankIFSCCode =  emp_Details[0]['BankIFSCCode']
    BankBranch  = emp_Details[0]['BankBranch']

    # print("the month no. is here::", month_no)
    # print("the year no. is here::", year)
      
    # Attendance = Attendance_Data.objects.filter(
    Attendance = SalaryAttendance.objects.filter(
        EmployeeCode=EmployeeCode,
        OrganizationID=OrganizationID,
        IsDelete = False,
        Date__month=month_no,
        Date__year=year
    ).order_by('Date')

    EmpName = emp_Details[0]['EmpName'] or ''
    Gender = emp_Details[0]['Gender'] or ''

    if fixed_basic: #and fixed_HRA:
        gross_salary = fixed_basic + fixed_HRA + Conveyance_Allowance + CCA + Other_Allowance

        try:
            att_lock = AttendanceLock.objects.filter(EmployeeCode = EmployeeCode,month = month_no,year= year,IsDelete = False,OrganizationID = OrganizationID).first() 
            total_no_Days_in_month = int(att_lock.total_no_Days_in_month)
            print(f"total_no_Days_in_month == {total_no_Days_in_month}")
            IsLock  = att_lock.IsLock
            # print(IsLock)
            # print(total_no_Days_in_month)
        except AttendanceLock.DoesNotExist:
            IsLock = None
            messages.error(request,f'Attendance is not Locked for {EmpName}')
            return redirect(reverse('Employees_Payroll_List')+f'?year={year}&month_no={month_no}')
        
            # return redirect(reverse('Daily_Attendace') + f'?Q_Date={Q_Date}') 
        
        no_of_days = int(att_lock.PaiDays)
        # print(f"no_of_days== {no_of_days}")
        no_of_absent = total_no_Days_in_month - no_of_days

        Earned_Basic = float(((fixed_basic / total_no_Days_in_month) * no_of_days))
        Earned_CCA=0
        if CCA>0:
            Earned_CCA = float(((CCA / total_no_Days_in_month) * no_of_days))
        
        Earned_Other_Allowance=0
        if Other_Allowance>0:
            Earned_Other_Allowance = float(((Other_Allowance / total_no_Days_in_month) * no_of_days))   
 

        Earned_HRA = float(((fixed_HRA / total_no_Days_in_month) * no_of_days))

        print("the stored arrear is here", Stored_Arrear)
        
        Front_Total_Earning = Earned_Basic + Earned_HRA + float(Stored_Arrear)
        Earned_Total_Allowance  =  Conveyance_Allowance + Earned_CCA + Earned_Other_Allowance
        Total_Earning  =  Front_Total_Earning + Earned_Total_Allowance

        # print(f"Total_Earning allowences::=:{Total_Earning}")

        # Front_Total_Earning_Demo = Earned_Basic + Earned_HRA + int(Stored_Arrear)
        # Earned_Total_Allowance_Demo  =  Conveyance_Allowance + Earned_CCA + Earned_Other_Allowance
        # Total_Earning_Demo  =  Front_Total_Earning_Demo + Earned_Total_Allowance_Demo

        # print("-----------------------------------------------------------------------")
        # print(f"{Front_Total_Earning} = {Earned_Basic} + {Earned_HRA} + {Stored_Arrear}")
        # print(f"{Earned_Total_Allowance} = {Conveyance_Allowance} + {Earned_CCA} + {Earned_Other_Allowance}")
        # print(f"{Total_Earning} = {Front_Total_Earning} + {Earned_Total_Allowance}")
        # print(f"Front_Total_Earning {Earned_Basic} + {Earned_HRA} + {int(Stored_Arrear)}  is here == {Front_Total_Earning}")
        # print(f"Earned_Total_Allowance {Conveyance_Allowance} + {Earned_CCA} + {Earned_Other_Allowance}  is here == {Earned_Total_Allowance}")
        # print(f"Total_Earning Demo {Front_Total_Earning} + {Earned_Total_Allowance} is here ==  {Total_Earning}")
        # print("-----------------------------------------------------------------------")



        # print("the new esic is here::", Front_Total_Earning)

        IsESICCalculate= OrgConfig.IsESICCalculate

        if OrgConfig.IsESICCalculate==False:
            ESIC = 0
            CompanyContributionToESIC=0
        elif gross_salary > 21000 or (OrgConfig is not None and OrgConfig.IsESICCalculate==False):
            ESIC = 0
            CompanyContributionToESIC=0
        else:    
            ESIC = int(((Front_Total_Earning/100)*0.75))
            CompanyContributionToESIC=int(((Front_Total_Earning/100)*3.25))

        EPFO=0
        if OrganizationID == '1401':
            PFCal = Earned_Basic+ Earned_CCA + Earned_Other_Allowance + Conveyance_Allowance + int(Stored_Arrear)
        elif OrganizationID == '1101':
            PFCal = Earned_Basic
        else:
            PFCal = Earned_Basic+ Earned_CCA + Earned_Other_Allowance + Conveyance_Allowance


        # print("The total pf cal is here::", PFCal)

        if PFCal > 15000:
            EPFO = EmployeePF
            # print("iam here ------ ::")
        else:
            # print("the EPFO is here ------ ::", EPFO)
            if EmployeePF>0:
                EPFO =  int(((PFCal/100)*12))
        
        # EPFO =  int(((PFCal/100)*12))

        # print("the EmployeePF is here::", EmployeePF)
        # print("the PFCalPFO is here::", PFCal)

        # print(f"the calculated is here:: {ESIC}")
        # print("The total earning is here:==:", Total_Earning)

        
        if EPFO>1800:
            EPFO=1800   

        # EPFO=0
        if EmployeePF>0:
            if fixed_basic > 15000:
                EmployeePF = 1950
            else: 
                EmployeePF = int(((Earned_Basic/100)*13))
                
        if OrganizationID=='2020':
            if Gender=="Female":
                # print("the Earned_Basic is already", Earned_Basic)
                if Total_Earning <25000:
                    PT=0
                elif Total_Earning>=25000:
                    if month_no==2:
                        PT=300
                    else:
                        PT=200
                # print("the PT is already", PT)
            
            else: 
                if Total_Earning<=7500:
                    PT=0
                elif Total_Earning>7500 and Total_Earning<=10000:
                    PT=175
                elif Total_Earning>10000:
                    if month_no==2:
                        PT=300
                    else:
                        PT=200
            
        elif OrganizationID=='1401':
                # print("-------------------------------------------")
                        
                # print("The total earning is here:==:", Total_Earning)
                if Total_Earning<=18750:
                    PT=0
                elif Total_Earning>18750 and Total_Earning<=25000:
                    PT=125
                elif Total_Earning>25000 and Total_Earning<=33333:
                    PT=167
                elif Total_Earning>33333:
                    if month_no==2:
                        PT=212
                    else:
                        PT=208
                
                print("the pt amount is here", PT)
            
                
        # elif OrganizationID=='1101' or OrganizationID=='1501' or OrganizationID=='501' or OrganizationID=='601' or OrganizationID=='601':
        elif OrganizationID == '1101':
            Demo_Earning = Front_Total_Earning + Earned_Total_Allowance
            if Demo_Earning>=10000:
                PT=200
            else:
                PT=0

        elif OrganizationID in ['1101', '501', '601', '20180612060935']:
            if Front_Total_Earning>=12000:
                PT=200
            else:
                PT=0

            print("the total earning is here::", Front_Total_Earning)

        elif OrganizationID == '2100':
            if fixed_Gross_Per_Annm <= 300000:
                PT = 0/12
            elif 300001 <= fixed_Gross_Per_Annm <= 500000:
                PT = 1200/12
            elif 500001 <= fixed_Gross_Per_Annm <= 800000:
                PT = 1800/12
            elif 800001 <= fixed_Gross_Per_Annm <= 1000000:
                PT = 2100/12
            else:  # salary >= 1000001
                PT = 2500/12

        elif OrganizationID == '1501':
            if Total_Earning>=12000:
                PT=200
            else:
                PT=0


        elif Earned_Basic<=10000:
            PT=0

                
        TotalCompanyContribution =	 CompanyContributionToESIC + EmployeePF
        Deduction = EPFO + ESIC + PT + Meals + Accommodation
        Front_Net_salary = int((Front_Total_Earning - Deduction))

        Front_Total_Deduction = int(Deduction)
        # Earned_Total_Allowance  =  Conveyance_Allowance + Earned_CCA + Earned_Other_Allowance
        # Total_Earning  =  Front_Total_Earning + Earned_Total_Allowance
        Net_salary = Front_Net_salary + Earned_Total_Allowance

        # print("The Front_Net_salary is here::", Front_Net_salary)
        # print("The Net_salary is here::", (Front_Net_salary - Earned_Total_Allowance))
        
        # if OrganizationID==1001:
       
        grid_data =  {
            'gross_salary':gross_salary,
            'total_no_Days_in_month':total_no_Days_in_month,
            'no_of_days':no_of_days,
            'no_of_absent':no_of_absent,
            'Total_Earning':'{:.1f}'.format(Total_Earning),
            'Earned_Basic':'{:.1f}'.format(Earned_Basic),
            'Earned_HRA':'{:.1f}'.format(Earned_HRA),
            'Earned_Total_Allowance':'{:.1f}'.format(Earned_Total_Allowance),

            'ESIC':'{:.1f}'.format(ESIC),
            'EPFO':'{:.1f}'.format(EPFO),
            'Total_Deduction':'{:.1f}'.format(Front_Total_Deduction),
            'Net_salary':'{:.1f}'.format(Net_salary),
            'PT':'{:.1f}'.format(PT),
        }

        Leave_Balance = Emp_Leave_Balance_Master.objects.filter(OrganizationID=OrganizationID,IsDelete =False,Emp_code=EmployeeCode)

        with transaction.atomic():
            if request.method == "POST":
                G_employee_code = request.POST['G_employee_code']
                G_month = request.POST['G_month']
                G_year = request.POST['G_year']
                PayMode = request.POST['PayMode']

                Arrear = request.POST['Arrear'] or 0
                RewardIncentive	 =   request.POST['RewardIncentive'] or 0
                # print("Earned_Total_Allowance",Earned_Total_Allowance)
                # Earned_Total_AllowanceC = request.POST['Earned_Total_AllowanceC'] or 0

                AdvanceLoan	=  request.POST['AdvanceLoan'] or 0
                TaxDeduction	=   request.POST['TaxDeduction'] or 0
                OtherDeduction	=    request.POST['OtherDeduction'] or 0 
                
                Total_Earning =   request.POST['TotalEarnings'] or 0
                Total_Deduction =  request.POST['Total_Deduction'] or 0
                Net_salary = float(request.POST['NetPay'] or 0)
                Net_salary_In_Words = number_to_words(int(Net_salary))

                FrontESIC =  request.POST['ESIC'] or 0


                if ID != '0':        
                    salary_slip = get_object_or_404(
                        SalarySlip,
                        id=ID,
                        IsDelete=False,
                        OrganizationID=OrganizationID
                    )

                    # print("we are here")

                    # update fields
                    salary_slip.month = G_month
                    salary_slip.year = G_year
                    salary_slip.month_name = month_name
                    salary_slip.PayMode = PayMode
                    salary_slip.Arrear = Arrear
                    salary_slip.RewardIncentive = RewardIncentive
                    salary_slip.AdvanceLoan = AdvanceLoan
                    salary_slip.TaxDeduction = TaxDeduction
                    salary_slip.OtherDeduction = OtherDeduction
                    salary_slip.Total_Earning = Total_Earning
                    salary_slip.Total_Deduction = Total_Deduction
                    salary_slip.Net_salary = Net_salary
                    salary_slip.Net_salary_In_Words = Net_salary_In_Words
                    salary_slip.ESIC = FrontESIC
                    salary_slip.EPFO = EPFO
                    salary_slip.PT = PT
                    salary_slip.Meals = Meals
                    salary_slip.Accommodation = Accommodation
                    salary_slip.EmployeePF = EmployeePF
                    salary_slip.CompanyContributionToESIC = CompanyContributionToESIC
                    salary_slip.TotalCompanyContribution = TotalCompanyContribution
                    salary_slip.CTC = Front_Total_Earning + TotalCompanyContribution
                    salary_slip.Earned_Total_Allowance =  Earned_Total_Allowance    
                    salary_slip.ModifyBy = UserID
                    salary_slip.IsLocked = True
                    salary_slip.save()

                    messages.success(request, f"Salary slip updated for {EmpName}")
                else:                        
                    salary_slip = SalarySlip.objects.create(
                        EmployeeCode = G_employee_code, 
                        Emp_Name = EmpName ,
                        
                        month =  G_month,
                        year =  G_year,
                        month_name = month_name, 
                        
                        generated = True ,
                        Desingation =  emp_Details[0]['Designation'],
                        Department =  emp_Details[0]['Department'],
                        DOJ =  emp_Details[0]['DateofJoiningDate'],
                        DOB =  formatted_date_DOB,
                        
                        total_no_Days_in_month = total_no_Days_in_month,  
                        no_of_days =   no_of_days,
                        no_of_absent = no_of_absent,
                        
                        fixed_basic =   fixed_basic,
                        fixed_HRA =  fixed_HRA,
                        ConveyanceAllowance = Conveyance_Allowance,
                        CCA	 = CCA,
                        OtherAllowance = Other_Allowance,
                        gross_salary =   gross_salary,
                        
                        
                        Earned_Basic =   Earned_Basic,
                        Earned_HRA =   Earned_HRA,
                        Arrear = Arrear  ,
                        RewardIncentive	 =  RewardIncentive,
                        Earned_Total_Allowance =  Earned_Total_Allowance,    
                        
                        ESIC =   FrontESIC,
                        EPFO =   EPFO,
                        PT	= PT,
                        Meals	= Meals,
                        Accommodation	= Accommodation,
                        AdvanceLoan	= AdvanceLoan,
                        TaxDeduction	= TaxDeduction,
                        OtherDeduction	= OtherDeduction,


                        EmployeePF = 	EmployeePF,
                        CompanyContributionToESIC =	CompanyContributionToESIC,
                        TotalCompanyContribution =	TotalCompanyContribution,
                        CTC = Front_Total_Earning + TotalCompanyContribution, 

                        
                        Total_Earning =   Total_Earning,
                        Total_Deduction =  Total_Deduction,
                        Net_salary =   Net_salary,
                        Net_salary_In_Words =   Net_salary_In_Words,
                        
                        ProvidentFundNumber =emp_Details[0]['ProvidentFundNumber'],
                        ESINumber =emp_Details[0]['ESINumber'],
                        BankAccountNumber = emp_Details[0]['BankAccountNumber'],
                        PayMode =  PayMode,
                        BankName  =  BankName,
                        BankIFSCCode = BankIFSCCode,
                        BankBranch =  BankBranch,
                                                            
                        OrganizationID =  OrganizationID,
                        CreatedBy =  UserID,
                        IsLocked = True
                    )
            
                    messages.success(request,f"Salary slip is created for {EmpName}")
                return redirect(reverse('Generate_Salary_Slip') + f'?year={G_year}&month_no={G_month}')
    else:
        messages.success(request,f"Salary details  of {EmpName} is not present")
            
        return redirect('Generate_Salary_Slip')  
    
    # print("entire attendece data:::", Attendance)


    # Convert list of dicts -> dict with Title as key
    emp_sal_Det_demo = {
        item['Title'].strip(): item for item in emp_sal_Det
    }

    # print("emp_sal_Det_demo::", emp_sal_Det_demo)
        
    # sal_obj_data = True 
    context = {
        'sal_obj':sal_obj,
        'org_Details':org_Details,
        "PT":PT,
        "IsESICCalculate":IsESICCalculate,
        'emp_sal_Det':emp_sal_Det,
        'emp_sal_Det_demo':emp_sal_Det_demo,
        'emp_Details':emp_Details,
        'month_no':month_no,
        'month':month_name,
        'year':year,
        'grid_data':grid_data,
        'Leave_Balance':Leave_Balance, 
        'attendance_data':Attendance,
        'IsLock':IsLock,
        "IsLocked_Slip": sal_obj.IsLocked if sal_obj else False,
        # "IsLocked_Slip":sal_obj_data,
        'Slip_ID':ID,
    }
    return render(request, "EMP_PAY/Salary/View_Emmployee_Salary_Details.html", context)










# -------------------------------------------------------------------- practice
from django.http import JsonResponse
from django.views.decorators.csrf import csrf_exempt

@csrf_exempt
def recalc_salary(request):
    """
    Accept POST with updated fields and return recalculated salary values
    """
    if request.method == "POST":
        emp_code = request.POST.get('G_employee_code')
        month = int(request.POST.get('G_month'))
        year = int(request.POST.get('G_year'))
        arrear = float(request.POST.get('Arrear', 0))
        reward = float(request.POST.get('RewardIncentive', 0))
        advance = float(request.POST.get('AdvanceLoan', 0))
        tax = float(request.POST.get('TaxDeduction', 0))
        other_ded = float(request.POST.get('OtherDeduction', 0))
        Form_ESIC = float(request.POST.get('ESIC', 0))
        Form_EPFO = float(request.POST.get('EPFO', 0))
        Form_Accommodation = float(request.POST.get('Accommodation', 0))
        Form_Meals = float(request.POST.get('Meals', 0))
        Form_PT = float(request.POST.get('PT', 0))


        print("----------------------------------------------------------------------")
        print("The Esic Value is here::", Form_ESIC)
        print("The Form_EPFO Value is here::", Form_EPFO)
        print("The Form_Accommodation Value is here::", Form_Accommodation)
        print("The Form_Meals Value is here::", Form_Meals)
        print("The Form_PT Value is here::", Form_PT)
        print("----------------------------------------------------------------------")


        print("The type of element (emp_code):", emp_code)
        print("The type of element (arrear):", arrear)
        print("The type of element (reward):", reward)
        print("The type of element (advance):", advance)
        print("The type of element (tax):", tax)
        print("The type of element (other_ded):", other_ded)

        
        OrganizationID = request.session.get("OrganizationID")
        
        # fetch employee salary details (from API or DB)
        hotelapitoken = MasterAttribute.HotelAPIkeyToken
        headers = {'hotel-api-token': hotelapitoken}
        api_url = f"http://hotelops.in/api/PyAPI/HREmployeeSalary?EmpCode={emp_code}&OrganizationID={OrganizationID}"
        
        try:
            response = requests.get(api_url, headers=headers)
            response.raise_for_status()
            emp_sal_Det = response.json()
        except:
            return JsonResponse({"error": "Unable to fetch salary data"}, status=400)
        
        # get attendance
        att_lock = AttendanceLock.objects.filter(EmployeeCode=emp_code, month=month, year=year, IsDelete=False, OrganizationID=OrganizationID).first()
        if not att_lock:
            return JsonResponse({"error": "Attendance not locked"}, status=400)
        
        total_days = float(att_lock.total_no_Days_in_month)
        paid_days = float(att_lock.PaiDays)
        
        # fixed_basic = get_salary_component(emp_sal_Det, "Basic") or 0
        # fixed_HRA = get_salary_component(emp_sal_Det, "HRA") or 0
        # CCA = get_salary_component(emp_sal_Det, "Conveyance Allowance") or 0
        # Other_Allowance = get_salary_component(emp_sal_Det, "Other Allowance") or 0

        print(f"{emp_sal_Det} ----------- Is Here:::::::::::::::::::::::::::::::::::::::::::::::::::::")


        fixed_basic_component = get_salary_component(emp_sal_Det, "Basic")
        fixed_HRA_component = get_salary_component(emp_sal_Det, "HRA")
        # Conveyance_Allowance_component = 0
        CCA_component = get_salary_component(emp_sal_Det, "Conveyance Allowance")
        Other_Allowance_component = get_salary_component(emp_sal_Det, "Other Allowance")
        PT_component = get_salary_component(emp_sal_Det, "PT")
        Meals_component = get_salary_component(emp_sal_Det, "Meals")
        Accommodation_component = get_salary_component(emp_sal_Det, "Accommodation")
        TTC_component = get_salary_component(emp_sal_Det, "Total Company Contribution (C)")
        EmployeePF_component = get_salary_component(emp_sal_Det, "Employee PF @12% (Basic)")
        ESIC_component = get_salary_component(emp_sal_Det, "ESIC @ 0.75%")

        Conveyance_Allowance  = 0

        fixed_basic = fixed_basic_component if fixed_basic_component else 0.0      
        fixed_HRA = fixed_HRA_component if fixed_HRA_component else 0.0      
        CCA = CCA_component if CCA_component else 0.0     
        Other_Allowance = Other_Allowance_component if Other_Allowance_component else 0.0      
        PT = PT_component if PT_component else 0.0      
        Meals = Meals_component if Meals_component else 0.0      
        Accommodation = Accommodation_component if Accommodation_component else 0.0      
        TTC = TTC_component if TTC_component else 0.0     
        EmployeePF = EmployeePF_component if EmployeePF_component else 0.0      
        ESIC = ESIC_component if ESIC_component else 0.0      

        
        
        print("The type of element (total_days):", total_days)
        print("The type of element (paid_days):", paid_days)
        print(f"The plus of  (paid_days)  - (total_days) :{total_days} - {paid_days}")
        # print("The type of element (fixed_basic):", type(fixed_basic))
        # Recalculate salary
        Earned_Basic = (fixed_basic / total_days) * paid_days
        Earned_HRA = (fixed_HRA / total_days) * paid_days
        Earned_CCA = (CCA / total_days) * paid_days
        Earned_Other = (Other_Allowance / total_days) * paid_days
        
        Total_Earning = Earned_Basic + Earned_HRA + Earned_CCA + Earned_Other + arrear + reward
        Total_Deduction = advance + tax + other_ded + PT + Meals + Accommodation + EmployeePF + ESIC
        Net_salary = Total_Earning - Total_Deduction
        
        # data = {
        #     "Earned_Basic": round(Earned_Basic, 2),
        #     "Earned_HRA": round(Earned_HRA, 2),
        #     "Earned_CCA": round(Earned_CCA, 2),
        #     "Earned_Other_Allowance": round(Earned_Other, 2),
        #     "Total_Earning": round(Total_Earning, 2),
        #     "Total_Deduction": round(Total_Deduction, 2),
        #     "Net_salary": round(Net_salary, 2)
        # }
        

        
        try:
            data = {
                "Earned_Basic": round(Earned_Basic, 2),
                "Earned_HRA": round(Earned_HRA, 2),
                "Earned_CCA": round(Earned_CCA, 2),
                "Earned_Other_Allowance": round(Earned_Other, 2),
                "Total_Earning": round(Total_Earning, 2),
                "Total_Deduction": round(Total_Deduction, 2),
                "Net_salary": round(Net_salary, 2)
            }

            return JsonResponse(data)
        except Exception as e:
            return JsonResponse({"status": "error", "message": str(e)}, status=500)
    else:
        return JsonResponse({"error": "Only POST allowed"}, status=400)
    


# -------------------------------------------------------------------- Form Handle


def Calculate_PT(request, OrganizationID, Earned_Basic, Total_Earning, fixed_Gross_Per_Annm, Front_Total_Earning, Gender, month_no):
    PT = 0  # default

    if OrganizationID == '2020':
        if Gender == "Female":
            if Total_Earning < 25000:
                PT = 0
            else:
                PT = 300 if month_no == 2 else 200
        else:
            if Total_Earning <= 7500:
                PT = 0
            elif Total_Earning <= 10000:
                PT = 175
            else:
                PT = 300 if month_no == 2 else 200

    elif OrganizationID == '1401':
        if Total_Earning <= 18750:
            PT = 0
        elif Total_Earning <= 25000:
            PT = 125
        elif Total_Earning <= 33333:
            PT = 167
        else:
            PT = 212 if month_no == 2 else 208

    elif OrganizationID in ['1101','601','20180612060935']:
        PT = 200 if Total_Earning >= 12000 else 0

    elif OrganizationID == '501':
        PT = 200 if Front_Total_Earning >= 12000 else 0

    elif OrganizationID == '2100':
        if fixed_Gross_Per_Annm <= 300000:
            PT = 0
        elif fixed_Gross_Per_Annm <= 500000:
            PT = 1200 // 12
        elif fixed_Gross_Per_Annm <= 800000:
            PT = 1800 // 12
        elif fixed_Gross_Per_Annm <= 1000000:
            PT = 2100 // 12
        else:
            PT = 2500 // 12

    elif OrganizationID == '1501':
        PT = 200 if Total_Earning >= 12000 else 0

    if Earned_Basic <= 10000:
        PT = 0

    return PT



# -------------------------------------------------------------------- practice


@transaction.atomic
def Unlock_Employee_Salary_Details(request, slip_id):
    slip = get_object_or_404(SalarySlip, id=slip_id, IsDelete=False)
    slip.IsLocked = False
    slip.save(update_fields=["IsLocked"])
    messages.success(request, f"Salary slip for {slip.Emp_Name} unlocked successfully")
    return redirect(request.META.get("HTTP_REFERER", "Employees_Payroll_List"))



import locale
@transaction.atomic
def Generate_Salary_Slip(request):
    if 'OrganizationID' not in request.session:
        return redirect(MasterAttribute.Host)
    else:
        print("Show Page Session")
   
    OrganizationID = request.session["OrganizationID"]
    UserID = str(request.session["UserID"])
    org_Details  =  OrganizationMaster.objects.get(OrganizationID= OrganizationID)
    
    hotelapitoken = MasterAttribute.HotelAPIkeyToken
    headers = {
    'hotel-api-token': hotelapitoken  
    }
    
    org_url = f"http://hotelops.in/API/PyAPI/OrganizationListSelect?OrganizationID={OrganizationID}"

    try:
        response = requests.get(org_url, headers=headers)
        response.raise_for_status()  
        memOrg = response.json()
    except requests.exceptions.RequestException as e:
        print(f"Error occurred: {e}")
     
    
    
    with transaction.atomic():
        Hr = request.GET.get('HR')
        if Hr:
            EmployeeCode = request.GET.get('Emp')
            
            api_ur = "http://hotelops.in/api/PyAPI/HREmployeeDataEmpCodeSelect?EmpCode="+str(EmployeeCode)+"&OID="+str(OrganizationID)

            try:
                response = requests.get(api_ur, headers=headers)
                response.raise_for_status()  
                emp_Details = response.json()
            
            except requests.exceptions.RequestException as e:
                print(f"Error occurred: {e}")
            
            emp_name = emp_Details[0]['EmpName']        
        else:
            EmployeeCode = request.session["EmployeeCode"]
            api_ur = "http://hotelops.in/api/PyAPI/HREmployeeDataEmpCodeSelect?EmpCode="+str(EmployeeCode)+"&OID="+str(OrganizationID)

            try:
                response = requests.get(api_ur, headers=headers)
                response.raise_for_status()  
                emp_Details = response.json()
            
            except requests.exceptions.RequestException as e:
                print(f"Error occurred: {e}")

            
        
        
        year = request.GET.get('year')
        month_no =  request.GET.get('month_no')
        month_name = calendar.month_name[int(month_no)]
        
        try:

            salary = SalarySlip.objects.get(OrganizationID=OrganizationID,IsDelete=False,
                                        EmployeeCode = EmployeeCode,month=month_no ,year =year,
                                        generated = True,
                                        HrVerify = True,
                                        FcVerify =  True
                                        )
           
           
            # locale.setlocale(locale.LC_NUMERIC, 'en_IN')


            # import locale
            try:
                locale.setlocale(locale.LC_ALL, "en_IN")   
            except locale.Error:
                locale.setlocale(locale.LC_ALL, "English_India.1252")  

   
          
            fixed_basic = locale.format_string("%.2f", float(salary.fixed_basic), grouping=True)
            fixed_HRA = locale.format_string("%.2f", float(salary.fixed_HRA), grouping=True)
            ConveyanceAllowance = locale.format_string("%.2f", float(salary.ConveyanceAllowance), grouping=True)
            CCA = locale.format_string("%.2f", float(salary.CCA), grouping=True)
            OtherAllowance = locale.format_string("%.2f", float(salary.OtherAllowance), grouping=True)
            gross_salary = locale.format_string("%.2f", float(salary.gross_salary), grouping=True)

            Earned_Basic = locale.format_string("%.2f", float(salary.Earned_Basic), grouping=True)
            Earned_HRA = locale.format_string("%.2f", float(salary.Earned_HRA), grouping=True)
            Arrear = locale.format_string("%.2f", float(salary.Arrear), grouping=True)
            RewardIncentive = locale.format_string("%.2f", float(salary.RewardIncentive), grouping=True)
            Total_Deduction = locale.format_string("%.2f", float(salary.Total_Deduction), grouping=True)
            Net_salary = locale.format_string("%.2f", float(salary.Net_salary), grouping=True)
            
            Total_Earning = locale.format_string("%.2f", float(salary.Total_Earning), grouping=True)



            ESIC = locale.format_string("%.2f", float(salary.ESIC), grouping=True)
            EPFO = locale.format_string("%.2f", float(salary.EPFO), grouping=True)
            PT = locale.format_string("%.2f", float(salary.PT), grouping=True)
            Meals = locale.format_string("%.2f", float(salary.Meals), grouping=True)
            Accommodation = locale.format_string("%.2f", float(salary.Accommodation), grouping=True)
            AdvanceLoan = locale.format_string("%.2f", float(salary.AdvanceLoan), grouping=True)
            TaxDeduction = locale.format_string("%.2f", float(salary.TaxDeduction), grouping=True)
            OtherDeduction = locale.format_string("%.2f", float(salary.OtherDeduction), grouping=True)
            Total_Deduction = locale.format_string("%.2f", float(salary.Total_Deduction), grouping=True)
            if salary.Earned_Total_Allowance is not None:
                Earned_Total_Allowance = locale.format_string("%.2f", float(salary.Earned_Total_Allowance), grouping=True)
            else:
                Earned_Total_Allowance = "0.00" 

            formatted = {
                    'fixed_basic' : fixed_basic ,
                    'fixed_HRA' : fixed_HRA ,
                    'ConveyanceAllowance' : ConveyanceAllowance,
                     'CCA' : CCA,
                    'OtherAllowance' : OtherAllowance,
                    'gross_salary' : gross_salary,
                    'Earned_Basic' : Earned_Basic,
                    'Earned_HRA' : Earned_HRA,
                    'Arrear' :  Arrear,
                    'RewardIncentive' :  RewardIncentive,
                    'Total_Deduction' :  Total_Deduction,
                    'Net_salary' :  Net_salary,
                    'ESIC' :  ESIC,
                    'EPFO' : EPFO,
                    'PT' : PT,
                    'Meals' : Meals ,
                    'Accommodation' : Accommodation ,
                    'AdvanceLoan' : AdvanceLoan,
                    'TaxDeduction' : TaxDeduction ,
                    'OtherDeduction' : OtherDeduction ,
                    'Total_Deduction' :  Total_Deduction,
                    'Total_Earning':Total_Earning,
                    'Earned_Total_Allowance':Earned_Total_Allowance

            }

             
        except SalarySlip.DoesNotExist:
            if Hr:
                messages.warning(request,f"No Salary Slip is present for {month_name} of {emp_name}")
                return redirect('Employees_Payroll_List')
            else:

                messages.warning(request,f"No Salary Slip is present for {month_name}")
            return redirect('Payroll_List')

        Leave_Balance = Emp_Leave_Balance_Master.objects.filter(OrganizationID=OrganizationID,IsDelete =False,Emp_code=EmployeeCode)
        template_path = "EMP_PAY/Salary/Generate_Salary_Slip.html"
        mydict = {'month':month_name,'year':year,'formatted' :formatted
                ,'salary':salary,'org_Details':org_Details,'Leave_Balance':Leave_Balance} 

        response = HttpResponse(content_type='application/pdf')
        response['Content-Disposition'] = 'filename="padp.pdf"'
        
        template = get_template(template_path)
        html = template.render(mydict)

        result = BytesIO()
        # pdf = pisa.pisaDocument(BytesIO(html.encode("ISO-8859-1")), result)
        pdf = pisa.pisaDocument(BytesIO(html.encode("UTF-8")), result)
        
        if not pdf.err:
            return HttpResponse(result.getvalue(), content_type='application/pdf')
        return None






def Generate_Salary_Slip_Emails(OrganizationID,UserID,EmployeeCode,year,month_no):
    # if 'OrganizationID' not in request.session:
    #     return redirect(MasterAttribute.Host)
    # else:
    #     print("Show Page Session")
   
    # OrganizationID = request.session["OrganizationID"]
    # UserID = str(request.session["UserID"])
    org_Details  =  OrganizationMaster.objects.get(OrganizationID= OrganizationID)
    
    hotelapitoken = MasterAttribute.HotelAPIkeyToken
    headers = {
    'hotel-api-token': hotelapitoken  
    }
    
    org_url = f"http://hotelops.in/API/PyAPI/OrganizationListSelect?OrganizationID={OrganizationID}"

    try:
        response = requests.get(org_url, headers=headers)
        response.raise_for_status()  
        memOrg = response.json()
    except requests.exceptions.RequestException as e:
        print(f"Error occurred: {e}")
     
    
    
    print(OrganizationID,UserID,EmployeeCode,year,month_no)
    
    with transaction.atomic():
            
        
            api_ur = "http://hotelops.in/api/PyAPI/HREmployeeDataEmpCodeSelect?EmpCode="+str(EmployeeCode)+"&OID="+str(OrganizationID)

            try:
                response = requests.get(api_ur, headers=headers)
                response.raise_for_status()  
                emp_Details = response.json()
            
            except requests.exceptions.RequestException as e:
                print(f"Error occurred: {e}")
                
            emp_name = emp_Details[0]['EmpName']        
            
                
            
            
            
            month_name = calendar.month_name[int(month_no)]
            
            try:

                salary = SalarySlip.objects.get(OrganizationID=OrganizationID,IsDelete=False,
                                            EmployeeCode = EmployeeCode,month=month_no ,year = year,
                                              generated = True,
                                            HrVerify = True,    
                                            FcVerify =  True
                                            )
                
                # locale.setlocale(locale.LC_NUMERIC, 'en_IN')

                # import locale
                try:
                    locale.setlocale(locale.LC_ALL, "en_IN")   # Linux/Mac
                except locale.Error:
                    locale.setlocale(locale.LC_ALL, "English_India.1252")  # Windows


                fixed_basic = locale.format_string("%.2f", float(salary.fixed_basic), grouping=True)
                fixed_HRA = locale.format_string("%.2f", float(salary.fixed_HRA), grouping=True)
                ConveyanceAllowance = locale.format_string("%.2f", float(salary.ConveyanceAllowance), grouping=True)
                CCA = locale.format_string("%.2f", float(salary.CCA), grouping=True)
                OtherAllowance = locale.format_string("%.2f", float(salary.OtherAllowance), grouping=True)
                gross_salary = locale.format_string("%.2f", float(salary.gross_salary), grouping=True)

                Earned_Basic = locale.format_string("%.2f", float(salary.Earned_Basic), grouping=True)
                Earned_HRA = locale.format_string("%.2f", float(salary.Earned_HRA), grouping=True)
                Arrear = locale.format_string("%.2f", float(salary.Arrear), grouping=True)
                RewardIncentive = locale.format_string("%.2f", float(salary.RewardIncentive), grouping=True)
                Total_Deduction = locale.format_string("%.2f", float(salary.Total_Deduction), grouping=True)
                Net_salary = locale.format_string("%.2f", float(salary.Net_salary), grouping=True)
                
                Total_Earning = locale.format_string("%.2f", float(salary.Total_Earning), grouping=True)

                ESIC = locale.format_string("%.2f", float(salary.ESIC), grouping=True)
                EPFO = locale.format_string("%.2f", float(salary.EPFO), grouping=True)
                PT = locale.format_string("%.2f", float(salary.PT), grouping=True)
                Meals = locale.format_string("%.2f", float(salary.Meals), grouping=True)
                Accommodation = locale.format_string("%.2f", float(salary.Accommodation), grouping=True)
                AdvanceLoan = locale.format_string("%.2f", float(salary.AdvanceLoan), grouping=True)
                TaxDeduction = locale.format_string("%.2f", float(salary.TaxDeduction), grouping=True)
                OtherDeduction = locale.format_string("%.2f", float(salary.OtherDeduction), grouping=True)
                Total_Deduction = locale.format_string("%.2f", float(salary.Total_Deduction), grouping=True)


                # fixed_basic = "{:,.2f}".format(float(salary.fixed_basic))
                # fixed_HRA = "{:,.2f}".format(float(salary.fixed_HRA))
                # ConveyanceAllowance = "{:,.2f}".format(float(salary.ConveyanceAllowance))
                # CCA = "{:,.2f}".format(float(salary.CCA))
                # OtherAllowance = "{:,.2f}".format(float(salary.OtherAllowance))
                # gross_salary = "{:,.2f}".format(float(salary.gross_salary))

                # Earned_Basic = "{:,.2f}".format(float(salary.Earned_Basic))
                # Earned_HRA = "{:,.2f}".format(float(salary.Earned_HRA))
                # Arrear = "{:,.2f}".format(float(salary.Arrear))
                # RewardIncentive = "{:,.2f}".format(float(salary.RewardIncentive))
                # Total_Deduction = "{:,.2f}".format(float(salary.Total_Deduction))
                # Net_salary = "{:,.2f}".format(float(salary.Net_salary))

                # Total_Earning = "{:,.2f}".format(float(salary.Total_Earning))

                # ESIC = "{:,.2f}".format(float(salary.ESIC))
                # EPFO = "{:,.2f}".format(float(salary.EPFO))
                # PT = "{:,.2f}".format(float(salary.PT))
                # Meals = "{:,.2f}".format(float(salary.Meals))
                # Accommodation = "{:,.2f}".format(float(salary.Accommodation))
                # AdvanceLoan = "{:,.2f}".format(float(salary.AdvanceLoan))
                # TaxDeduction = "{:,.2f}".format(float(salary.TaxDeduction))
                # OtherDeduction = "{:,.2f}".format(float(salary.OtherDeduction))
                # Total_Deduction = "{:,.2f}".format(float(salary.Total_Deduction))

                # if salary.Earned_Total_Allowance is not None:
                #     Earned_Total_Allowance = "{:,.2f}".format(float(salary.Earned_Total_Allowance))
                # else:
                #     Earned_Total_Allowance = "0.00"

                formatted = {
                        'fixed_basic' : fixed_basic ,
                        'fixed_HRA' : fixed_HRA ,
                        'ConveyanceAllowance' : ConveyanceAllowance,
                        'CCA' : CCA,
                        'OtherAllowance' : OtherAllowance,
                        'gross_salary' : gross_salary,
                        'Earned_Basic' : Earned_Basic,
                        'Earned_HRA' : Earned_HRA,
                        'Arrear' :  Arrear,
                        'RewardIncentive' :  RewardIncentive,
                        'Total_Deduction' :  Total_Deduction,
                        'Net_salary' :  Net_salary,
                        'ESIC' :  ESIC,
                        'EPFO' : EPFO,
                        'PT' : PT,
                        'Meals' : Meals ,
                        'Accommodation' : Accommodation ,
                        'AdvanceLoan' : AdvanceLoan,
                        'TaxDeduction' : TaxDeduction ,
                        'OtherDeduction' : OtherDeduction ,
                        'Total_Deduction' :  Total_Deduction,
                        'Total_Earning':Total_Earning

                }
                Leave_Balance = Emp_Leave_Balance_Master.objects.filter(OrganizationID=OrganizationID,IsDelete =False,Emp_code=EmployeeCode)
                template_path = "EMP_PAY/Salary/Generate_Salary_Slip.html"
                
                
                
                mydict = {'month':month_name,'year':year
                        ,'salary':salary,'org_Details':org_Details,'Leave_Balance':Leave_Balance,'formatted':formatted} 

                response = HttpResponse(content_type='application/pdf')
                response['Content-Disposition'] = 'filename="padp.pdf"'
                
                template = get_template(template_path)
                html = template.render(mydict)

                result = BytesIO()
                pdf = pisa.pisaDocument(BytesIO(html.encode("ISO-8859-1")), result)
                
                if not pdf.err:
                    return result.getvalue()  
                return None
                
            except SalarySlip.DoesNotExist:
                
                ErrorMessage = f"No Salary Slip is present for {month_name}"
                ErrorPage = "Generate_Salary_Slip_Emails"
                ErrorFucntion =  "Generate_Salary_Slip_Emails()"
                errorlog = PayrollErrorLog.objects.create(OrganizationID = OrganizationID,ErrorMessage = ErrorMessage,ErrorPage  = ErrorPage, ErrorFucntion = ErrorFucntion,CreatedBy= UserID)
            

        
            
            




from django.core.mail import EmailMessage
from django.template.loader import render_to_string
from django.conf import settings
def SalarySendEmails(request):
    if 'OrganizationID' not in request.session:
        return redirect(MasterAttribute.Host)
    else:
        print("Show Page Session")
   
    OrganizationID = request.session["OrganizationID"]
    UserID = str(request.session["UserID"])
    hotelapitoken = MasterAttribute.HotelAPIkeyToken
    headers = {
        'hotel-api-token': hotelapitoken  
    }
    api_url = "http://hotelops.in/API/PyAPI/HREmployeeList?OrganizationID="+str(OrganizationID)
    

    try:
        response = requests.get(api_url, headers=headers)
        response.raise_for_status() 
        emp_list = response.json()
     
    except requests.exceptions.RequestException as e:
        print(f"Error occurred: {e}")
    
    if request.method == "POST":
            year = request.POST.get('year')
            month_no =  request.POST.get('month_no')
            
            month_name = calendar.month_name[int(month_no)]
            for emp in emp_list:
            
                EmployeeCode = emp['EmployeeCode']
                # Remove if when live
                if  EmployeeCode  == '451000245':
                        emp_data_url  = "http://hotelops.in/api/PyAPI/HREmployeeDataEmpCodeSelect?EmpCode="+str(EmployeeCode)+"&OID="+str(OrganizationID)

                        try:
                            response = requests.get(emp_data_url, headers=headers)
                            response.raise_for_status()  
                            emp_Details = response.json()
                        
                        except requests.exceptions.RequestException as e:
                            print(f"Error occurred: {e}")
                        
                        emp_name = emp_Details[0]['EmpName']
                        # emp_email  = emp_Details[0]['EmailAddress']
                        # update when live
                        emp_email = ["darpananjanaynilehospitality@gmail.com" ]
                    
                        pdf_content = Generate_Salary_Slip_Emails(OrganizationID,UserID, EmployeeCode, year, month_no)
                      
                        subject =  f'Salary Slip of {month_name},{year}'
                        message = f'Dear {emp_name} Please Find attachement of your salary slip of {month_name},{year}'

                        try:
                            if pdf_content:
                                email = EmailMessage(subject, message, settings.EMAIL_HOST_USER, emp_email)
                                email.content_subtype = "html"
                                email.attach('attachment.pdf', pdf_content, 'application/pdf')
                                email.send()
                                email_obj = SalaryEmails.objects.create(OrganizationID = OrganizationID,CreatedBy =  UserID,
                                                                        EmpCode = EmployeeCode,
                                                                        email = emp_email)  
                        except requests.exceptions.RequestException as e:
                                        
                        
                                ErrorMessage = f"Error while sending Email {e}"
                                ErrorPage = "SendEmails"
                                ErrorFucntion =  "SendEmails()"
                                errorlog = PayrollErrorLog.objects.create(OrganizationID = OrganizationID,ErrorMessage = ErrorMessage,ErrorPage  = ErrorPage, ErrorFucntion = ErrorFucntion,CreatedBy =  UserID)
                    
                        messages.success(request,'Sended Successfully')        
                        return redirect('SalarySendEmails')


       
    today = datetime.today()
    CYear = today.year
    CMonth = today.month

    context = {'CYear':range(CYear,2020,-1),'CMonth':CMonth}
    return render(request, "EMP_PAY/Email/SalarySendEmails.html", context)
    


def number_to_words(number):
  
    units = ['Zero', 'One', 'Two', 'Three', 'Four', 'Five', 'Six', 'Seven', 'Eight', 'Nine',
             'Ten', 'Eleven', 'Twelve', 'Thirteen', 'Fourteen', 'Fifteen', 'Sixteen', 'Seventeen', 'Eighteen', 'Nineteen']
    
    tens = ['', '', 'Twenty', 'Thirty', 'Forty', 'Fifty', 'Sixty', 'Seventy', 'Eighty', 'Ninety']
    
   
    powers_of_ten = ['', 'Thousand', 'Million', 'Billion', 'Trillion', 'Quadrillion', 'Quintillion']

    
    def less_than_thousand(num):
        if num == 0:
            return ''
        elif num < 20:
            return units[num]
        elif num < 100:
            return tens[num // 10] + ('' if num % 10 == 0 else ' ' + units[num % 10])
        else:
            return units[num // 100] + ' Hundred' + ('' if num % 100 == 0 else ' and ' + less_than_thousand(num % 100))

   
    def convert_to_words(num):
        if num == 0:
            return 'Zero'
        words = ''
        for i in range(len(powers_of_ten)):
            if num % 1000 != 0:
                words = less_than_thousand(num % 1000) + ' ' + powers_of_ten[i] + ' ' + words
            num //= 1000
        return words.strip()

    return convert_to_words(number)


from django.db.models import Sum

    # Totals Aggregation
def safe_int(value):
    try:
        return int(float(value))
    except (ValueError, TypeError):
        return 0
    
    
# Salary List of Employees
def SalaryList(request):
    if 'OrganizationID' not in request.session:
        return redirect(MasterAttribute.Host)
    else:
        print("Show Page Session")
     
    OrganizationID = request.session["OrganizationID"]
    UserID = str(request.session["UserID"])
    UserType =   request.session["UserType"]
    Department_Name = request.session["Department_Name"]

   

    hotelapitoken = MasterAttribute.HotelAPIkeyToken
    headers = {
        'hotel-api-token': hotelapitoken  # Replace with your actual header key and value
    }
    api_url = "http://hotelops.in/API/PyAPI/OrganizationListSelect?OrganizationID=" + str(OrganizationID)
    memOrg=None
    try:
        response = requests.get(api_url, headers=headers)
        response.raise_for_status()  # Optional: Check for any HTTP errors
        memOrg = response.json()
    except requests.exceptions.RequestException as e:
        print(f"Error occurred: {e}")   

    I = request.GET.get('I',OrganizationID)
    
    year = request.GET.get('year')
    print("selected year is here::", year)

    if year:
        year = int(year)
    else:
        year = datetime.now().year

    month_no =  request.GET.get('month_no')
    if month_no:
        month_no = int(month_no)
    else:
        month_no = datetime.now().month  
    month_name = calendar.month_name[int(month_no)]    

    PayMode =request.GET.get('PayMode','')
    Status = request.GET.get('Status',0)

    

    if Status is None:
        VerifyStatus = None
    elif Status == '2':
        VerifyStatus = None  
    else:
        VerifyStatus = bool(int(Status))  
    
    
    if I == '':
        I= OrganizationID
    

    sal_list = SalarySlip.objects.filter(
        OrganizationID=I,
        IsDelete=False,
        month=month_no,
        year=year,
        generated=True
    ).order_by('EmployeeCode')

  
    if UserType.lower() == "hod":
        if Department_Name.lower() == "hr":
            if VerifyStatus == None:
                sal_list = sal_list
            else:
                sal_list = sal_list.filter(HrVerify=VerifyStatus)

        elif Department_Name.lower() == "finance":
                if VerifyStatus == None:
                    sal_list = sal_list.filter(HrVerify=True)
                else:    
                    sal_list = sal_list.filter(FcVerify=VerifyStatus,HrVerify=True)

    if PayMode:
        sal_list = sal_list.filter(PayMode=PayMode)
    
    totals = {
        'total_paid_days': sum(safe_int(s.no_of_days) for s in sal_list),
        'total_net_salary': sum(safe_int(s.Net_salary) for s in sal_list),
        'total_fixed_basic': sum(safe_int(s.fixed_basic) for s in sal_list),
        'total_fixed_HRA': sum(safe_int(s.fixed_HRA) for s in sal_list),
        'total_ConveyanceAllowance': sum(safe_int(s.ConveyanceAllowance) for s in sal_list),
        'total_CCA': sum(safe_int(s.CCA) for s in sal_list),
        'total_OtherAllowance': sum(safe_int(s.OtherAllowance) for s in sal_list),
        'total_gross_salary': sum(safe_int(s.gross_salary) for s in sal_list),
        'total_Earned_Basic': sum(safe_int(s.Earned_Basic) for s in sal_list),
        'total_Earned_HRA': sum(safe_int(s.Earned_HRA) for s in sal_list),
        'total_Arrear': sum(safe_int(s.Arrear) for s in sal_list),
        'total_RewardIncentive': sum(safe_int(s.RewardIncentive) for s in sal_list),
        'total_Total_Earning': sum(safe_int(s.Total_Earning) for s in sal_list),
        'total_ESIC': sum(safe_int(s.ESIC) for s in sal_list),
        'total_EPFO': sum(safe_int(s.EPFO) for s in sal_list),
        'total_PT': sum(safe_int(s.PT) for s in sal_list),
        'total_Meals': sum(safe_int(s.Meals) for s in sal_list),
        'total_Accommodation': sum(safe_int(s.Accommodation) for s in sal_list),
        'total_AdvanceLoan': sum(safe_int(s.AdvanceLoan) for s in sal_list),
        'total_TaxDeduction': sum(safe_int(s.TaxDeduction) for s in sal_list),
        'total_OtherDeduction': sum(safe_int(s.OtherDeduction) for s in sal_list),
        'total_Total_Deduction': sum(safe_int(s.Total_Deduction) for s in sal_list),
        'total_EmployeePF': sum(safe_int(s.EmployeePF) for s in sal_list),
        'total_CompanyContributionToESIC': sum(safe_int(s.CompanyContributionToESIC) for s in sal_list),
        'total_TotalCompanyContribution': sum(safe_int(s.TotalCompanyContribution) for s in sal_list),
        'total_CTC': sum(safe_int(s.CTC) for s in sal_list),
    }
        
    today = datetime.today()
    CYear = today.year
    CMonth = today.month    
    # print("year is here::", year)
    context = {
            'sal_list':sal_list,'I':I,'memOrg':memOrg,'PayMode':PayMode,
            'CYear':range(CYear,2020,-1),'CMonth':CMonth,
            'month_no':month_no,
            'month_name':month_name,
            'year':year,
            'Status':str(Status),
            'totals': totals,
        }
    return render(request,"EMP_PAY/Salary/SalaryList.html", context)


# def Salary_List_Pdf(request):
#     if 'OrganizationID' not in request.session:
#         return redirect(MasterAttribute.Host)

#     OrganizationID = request.session["OrganizationID"]
#     UserID = str(request.session["UserID"])
#     UserType =   request.session["UserType"]
#     Department_Name = request.session["Department_Name"]

#     # Get filter parameters
#     # Status = request.GET.get('Status', 0).strip()
#     month_no = request.GET.get('month_no', '').strip()
#     year = request.GET.get('year', '').strip()
#     # PayMode = request.GET.get('PayMode', '').strip()
#     Selected_OrganizationID = request.GET.get('I',OrganizationID)

#     Status = request.GET.get('Status', '0').strip() or '0'
#     PayMode = request.GET.get('PayMode', '').strip()


#     # print("Status",Status)
#     # print("month_no",month_no)
#     # print("year",year)
#     # print("PayMode",PayMode)
#     # print("Selected_OrganizationID",Selected_OrganizationID)



#     if Status is None:
#         VerifyStatus = None
#     elif Status == '2':
#         VerifyStatus = None  
#     else:
#         try:
#             VerifyStatus = bool(int(Status))
#         except ValueError:
#             VerifyStatus = None 
    
#     if Selected_OrganizationID == '':
#         Selected_OrganizationID= OrganizationID
    

#     # Base queryset
#     sal_list = SalarySlip.objects.filter(
#         OrganizationID=Selected_OrganizationID,
#         IsDelete=False,
#         month=month_no,
#         year=year,
#         generated=True
#     ).order_by('EmployeeCode')

      
#     if UserType.lower() == "hod":
#         if Department_Name.lower() == "hr":
#             if VerifyStatus == None:
#                 sal_list = sal_list
#             else:
#                 sal_list = sal_list.filter(HrVerify=VerifyStatus)

#         elif Department_Name.lower() == "finance":
#                 if VerifyStatus == None:
#                     sal_list = sal_list.filter(HrVerify=True)
                    
#                 else:    
           
#                     sal_list = sal_list.filter(FcVerify=VerifyStatus,HrVerify=True)


#     # Apply filters
#     if month_no:
#         sal_list = sal_list.filter(HrVerify=VerifyStatus)

#     if year:
#         sal_list = sal_list.filter(HrVerify=VerifyStatus)

#     if PayMode:
#         sal_list = sal_list.filter(HrVerify=VerifyStatus)


#     base_url = "https://hotelopsblob.blob.core.windows.net/hotelopslogos/"
#     organizations = OrganizationMaster.objects.filter(OrganizationID=3).first()
#     organization_logos = f"{base_url}{organizations.OrganizationLogo}" if organizations and organizations.OrganizationLogo else None

#     organizations = OrganizationMaster.objects.filter(OrganizationID=OrganizationID).first()
#     organization_logo = f"{base_url}{organizations.OrganizationLogo}" if organizations and organizations.OrganizationLogo else None
#     organization_logo = organizations.OrganizationName
#     # print(organization_logo)

#     formatted_sal_list = []

#     for sal in sal_list:
#         sal.DOJ_display = format_date_string(sal.DOJ)  # Add new formatted field
#         print("sal.DOJlay", sal.DOJ)
#         print("sal.DOJ_display", sal.DOJ_display)
#         formatted_sal_list.append(sal)

#     current_datetime = datetime.now().strftime('%d %B %Y %H:%M:%S')
#     context = {
#         'sal_list':formatted_sal_list,
#         'organization_logo': organization_logo,
#         'organization_logos':organization_logos,
#         'current_datetime':current_datetime,
#         'selected_month': datetime.now().month,
#     }

#     template_path = 'EMP_PAY/Salary/SalaryList_PDF.html'
#     # template_path = 'LMS/APOOVAL/EmpLeaveDataPdf.html'
#     # Employee_Payroll\templates\EMP_PAY\Salary\SalaryList_PDF.html
#     template = get_template(template_path)
#     html = template.render(context).encode("UTF-8")

#     response = HttpResponse(content_type='application/pdf')
#     response['Content-Disposition'] = f'attachment; filename="Salary_List{organization_logo}.pdf"'
#     # response['Content-Disposition'] = f'attachment; filename="{organization_logo}_{start_date}__To__{end_date}.pdf"'

#     pisa_status = pisa.CreatePDF(html, dest=response, encoding="UTF-8")

#     if pisa_status.err:
#         return HttpResponse("Error generating PDF", status=500)

#     return response

from django.http import HttpResponse, JsonResponse
from django.template.loader import get_template
from datetime import datetime
from xhtml2pdf import pisa
from django.db import connection



# def Salary_List_Pdf(request):
#     if 'OrganizationID' not in request.session:
#         return redirect(MasterAttribute.Host)

#     OrganizationID = request.session["OrganizationID"]
#     UserID = str(request.session["UserID"])
#     UserType = request.session["UserType"]
#     Department_Name = request.session["Department_Name"]

#     # Get filter parameters
#     month_no = request.GET.get('month_no', '').strip()
#     year = request.GET.get('year', '').strip()
#     Selected_OrganizationID = request.GET.get('I', OrganizationID)

#     if not Selected_OrganizationID:
#         Selected_OrganizationID = OrganizationID



#     base_url = "https://hotelopsblob.blob.core.windows.net/hotelopslogos/"

#     org_obj = OrganizationMaster.objects.filter(OrganizationID=3).first()
#     organization_logos = f"{base_url}{org_obj.OrganizationLogo}" if org_obj and org_obj.OrganizationLogo else None

#     org_obj = OrganizationMaster.objects.filter(OrganizationID=OrganizationID).first()
#     organization_logo = f"{base_url}{org_obj.OrganizationLogo}" if org_obj and org_obj.OrganizationLogo else None
#     organization_name = org_obj.OrganizationName if org_obj else ''

#     current_datetime = datetime.now().strftime('%d %B %Y %H:%M:%S')

#     context = {
#         # 'sal_list': raw_results,
#         'organization_logo': organization_name,
#         'organization_logos': organization_logos,
#         'current_datetime': current_datetime,
#         'selected_month': datetime.now().month,
#     }

#     template_path = 'EMP_PAY/Salary/SalaryList_PDF.html'
#     template = get_template(template_path)
#     html = template.render(context).encode("UTF-8")

#     response = HttpResponse(content_type='application/pdf')
#     response['Content-Disposition'] = f'attachment; filename="Salary_List_{organization_name}.pdf"'

#     pisa_status = pisa.CreatePDF(html, dest=response, encoding="UTF-8")

#     if pisa_status.err:
#         return HttpResponse("Error generating PDF", status=500)

#     return response

from django.db import connection
from django.template.loader import get_template
from django.http import HttpResponse
from xhtml2pdf import pisa
from datetime import datetime

# def Salary_List_Pdf(request):
#     if 'OrganizationID' not in request.session:
#         return redirect(MasterAttribute.Host)

#     OrganizationID = request.session["OrganizationID"]
#     month_no = request.GET.get('month_no', '').strip()
#     year = request.GET.get('year', '').strip()
#     Selected_OrganizationID = request.GET.get('I', OrganizationID)

#     # Execute stored procedure
#     with connection.cursor() as cursor:
#         cursor.execute("""
#             EXEC Employee_Payroll_SP_SalaryGenrate_PDF_Report_Select 
#                 @OrganizationID=%s, 
#                 @Month=%s, 
#                 @Year=%s
#         """, [Selected_OrganizationID, month_no, year])

#         columns = [col[0] for col in cursor.description]
#         results = [
#             dict(zip(columns, row))
#             for row in cursor.fetchall()
#         ]

#     # Optional: format any fields if needed, e.g. DOJ
#     for row in results:
#         doj = row.get('DOJ')
#         if doj:
#             row['DOJ_display'] = doj.strftime('%d-%m-%Y') if isinstance(doj, datetime) else doj

#     # Load logo
#     base_url = "https://hotelopsblob.blob.core.windows.net/hotelopslogos/"
#     organizations = OrganizationMaster.objects.filter(OrganizationID=3).first()
#     organization_logos = f"{base_url}{organizations.OrganizationLogo}" if organizations and organizations.OrganizationLogo else None

#     organizations = OrganizationMaster.objects.filter(OrganizationID=OrganizationID).first()
#     organization_logo = f"{base_url}{organizations.OrganizationLogo}" if organizations and organizations.OrganizationLogo else None
#     organization_name = organizations.OrganizationName if organizations else ""

#     current_datetime = datetime.now().strftime('%d %B %Y %H:%M:%S')
#     context = {
#         'sal_list': results,
#         'organization_logo': organization_name,
#         'organization_logos': organization_logos,
#         'current_datetime': current_datetime,
#         'selected_month': datetime.now().month,
#     }

#     template_path = 'EMP_PAY/Salary/SalaryList_PDF.html'
#     template = get_template(template_path)
#     html = template.render(context).encode("UTF-8")

#     response = HttpResponse(content_type='application/pdf')
#     response['Content-Disposition'] = f'attachment; filename="Salary_List_{organization_name}.pdf"'

#     pisa_status = pisa.CreatePDF(html, dest=response, encoding="UTF-8")

#     if pisa_status.err:
#         return HttpResponse("Error generating PDF", status=500)

#     return response


# def Salary_List_Pdf(request):
#     if 'OrganizationID' not in request.session:
#         return redirect(MasterAttribute.Host)

#     OrganizationID = request.session["OrganizationID"]
#     UserID = str(request.session["UserID"])
#     UserType = request.session["UserType"]
#     Department_Name = request.session["Department_Name"]

#     # Get filter parameters
#     month_no = request.GET.get('month_no', '').strip()
#     year = request.GET.get('year', '').strip()
#     Selected_OrganizationID = request.GET.get('I', OrganizationID)

#     print("Month-no::", type(month_no), month_no)
#     print("year-no::", type(year), year)
#     print("OrganizationID-no::", type(OrganizationID))
#     print("Selected_OrganizationID-no::", type(Selected_OrganizationID))
#     month_name = ''

#     if month_no.isdigit() and year.isdigit():
#         month_num = int(month_no)
#         year_num = int(year)
        
#         month_name = calendar.month_name[month_num]
#         _, month_days = calendar.monthrange(year_num, month_num)

#     print("Month-no::", type(month_no), month_no)
#     print("year-no::", type(year), year)

#     if not Selected_OrganizationID:
#         Selected_OrganizationID = OrganizationID

#     # Call stored procedure and fetch data
#     with connection.cursor() as cursor:
#         cursor.execute("""
#             EXEC Employee_Payroll_SP_SalaryGenrate_PDF_Report_Select 
#                 @OrganizationID=%s, 
#                 @Month=%s, 
#                 @Year=%s
#         """, [int(Selected_OrganizationID), int(month_no), int(year)])
        
#         columns = [col[0] for col in cursor.description]
#         sal_list = [dict(zip(columns, row)) for row in cursor.fetchall()]
    
#     # Format DOJ using the custom function
#     for row in sal_list:
#         doj = row.get('DOJ')
#         row['DOJ_display'] = format_date_string(doj)

#     # Organization logo and name setup
#     base_url = "https://hotelopsblob.blob.core.windows.net/hotelopslogos/"
#     org_obj = OrganizationMaster.objects.filter(OrganizationID=3).first()
#     organization_logos = f"{base_url}{org_obj.OrganizationLogo}" if org_obj and org_obj.OrganizationLogo else None

#     org_obj = OrganizationMaster.objects.filter(OrganizationID=OrganizationID).first()
#     organization_logo = f"{base_url}{org_obj.OrganizationLogo}" if org_obj and org_obj.OrganizationLogo else None
#     organization_name = org_obj.OrganizationName if org_obj else ''

#     current_datetime = datetime.now().strftime('%d %B %Y %H:%M:%S')

#     context = {
#         'sal_list': sal_list,
#         'month_days':month_days,
#         'month_name':month_name,
#         'year':year,
#         'organization_logo': organization_name,
#         'organization_logos': organization_logos,
#         'current_datetime': current_datetime,
#         'selected_month': datetime.now().month,
#     }

#     # Render to PDF
#     template_path = 'EMP_PAY/Salary/SalaryList_PDF.html'
#     template = get_template(template_path)
#     html = template.render(context).encode("UTF-8")

#     response = HttpResponse(content_type='application/pdf')
#     response['Content-Disposition'] = f'attachment; filename="Salary_List_{organization_name}.pdf"'

#     pisa_status = pisa.CreatePDF(html, dest=response, encoding="UTF-8")

#     if pisa_status.err:
#         return HttpResponse("Error generating PDF", status=500)

#     return response

from datetime import datetime






def Salary_List_Pdf(request):
    if 'OrganizationID' not in request.session:
        return redirect(MasterAttribute.Host)

    OrganizationID = request.session["OrganizationID"]
    UserID = str(request.session["UserID"])
    UserType = request.session["UserType"]
    Department_Name = request.session["Department_Name"]

    # Filters
    I = request.GET.get("I", OrganizationID)
    year = request.GET.get("year")
    month_no = request.GET.get("month_no")
    PayMode = request.GET.get("PayMode", "")
    Status = request.GET.get("Status", 0)

    # Defaults
    year = int(year) if year else datetime.now().year
    month_no = int(month_no) if month_no else datetime.now().month
    month_name = calendar.month_name[month_no]
    _, month_days = calendar.monthrange(year, month_no)

    # Status handling
    if Status is None or Status == "2":
        VerifyStatus = None
    else:
        VerifyStatus = bool(int(Status))

    if not I:
        I = OrganizationID

    # ORM query instead of stored procedure
    sal_list = SalarySlip.objects.filter(
        OrganizationID=I,
        IsDelete=False,
        month=month_no,
        year=year,
        generated=True
    ).order_by("EmployeeCode")

    if UserType.lower() == "hod":
        if Department_Name.lower() == "hr":
            if VerifyStatus is not None:
                sal_list = sal_list.filter(HrVerify=VerifyStatus)
        elif Department_Name.lower() == "finance":
            if VerifyStatus is None:
                sal_list = sal_list.filter(HrVerify=True)
            else:
                sal_list = sal_list.filter(FcVerify=VerifyStatus, HrVerify=True)

    if PayMode:
        sal_list = sal_list.filter(PayMode=PayMode)

    # Add computed fields (like in SalaryList)
    # for s in sal_list:
    #     total_earning = safe_int(s.Total_Earning)
    #     total_contrib = safe_int(s.TotalCompanyContribution)
    #     s.CTC = total_earning + total_contrib
    #     s.DOJ_display = format_date_string(s.DOJ)


    for s in sal_list:
        total_earning = safe_int(s.Total_Earning)
        EmployeePF = safe_int(s.EmployeePF)
        fixed_basic = safe_int(s.fixed_basic)
        Earned_Basic = safe_int(s.Earned_Basic)
        total_company_contribution = safe_int(s.TotalCompanyContribution)
        Totalctccalculated = total_earning + total_company_contribution
        s.CTC = total_earning + total_company_contribution
        s.DOJ_display = format_date_string(s.DOJ)


        Earned_Basic = safe_int(s.Earned_Basic)
        Earned_HRA = safe_int(s.Earned_HRA) 
        Earned_Total_Allowance = safe_int(s.Earned_Total_Allowance) 
        gross_salary = safe_int(s.gross_salary) 
        # print("the fixed_basic is here::", s.fixed_basic)
        s.Earned_Basic = Earned_Basic
        s.Earned_HRA = Earned_HRA
        s.Earned_Total_Allowance = Earned_Total_Allowance
        s.gross_salary = gross_salary

        
        Total_Earning = s.Total_Earning if s.Total_Earning else 0
        # Earned_HRA = s.Earned_HRA if s.Earned_HRA else 0 
        # s.Earned_HRA = safe_int(Earned_Basic)
        Total_PF_Data = float(Total_Earning) - float(Earned_HRA)


        # print("total earning is here:----:", Total_PF_Data)


        EmployeePF_data_demo = 0
        EmployeePF_data = 0
        if EmployeePF>0:
            if Total_PF_Data > 15000:
                EmployeePF_data = 1950
                # s.EmployeePF = EmployeePF_data
            else: 
                EmployeePF_data = int(((Total_PF_Data/100)*13))
                # s.EmployeePF = EmployeePF_data


        CompanyContributionToESIC_value = s.CompanyContributionToESIC if s.CompanyContributionToESIC else 0

        tcctype1 = float(CompanyContributionToESIC_value)
        employeedatatype1= float(EmployeePF_data)
        total = tcctype1 + employeedatatype1

        s.TotalCompanyContribution = total if total else 0
        # print(f"Total_Company_Contribution_Calc:==: {TotalCompanyContribution_value}")
        # print(f"Total_Company_Contribution_Calc:==: {tcctype1} -+- {employeedatatype1} == {total}")


    # Organization logo + name
    base_url = "https://hotelopsblob.blob.core.windows.net/hotelopslogos/"
    org_obj = OrganizationMaster.objects.filter(OrganizationID=OrganizationID).first()
    organization_logo = f"{base_url}{org_obj.OrganizationLogo}" if org_obj and org_obj.OrganizationLogo else None
    organization_name = org_obj.OrganizationName if org_obj else ""

    # Context for PDF
    context = {
        "sal_list": sal_list,
        "month_days": month_days,
        "month_name": month_name,
        "year": year,
        "organization_logo": organization_logo,
        "organization_name": organization_name,
        "current_datetime": datetime.now().strftime("%d %B %Y %H:%M:%S"),
        "selected_month": datetime.now().month,
    }

    # Render to PDF
    template_path = "EMP_PAY/Salary/SalaryList_PDF.html"
    template = get_template(template_path)
    html = template.render(context).encode("UTF-8")

    response = HttpResponse(content_type="application/pdf")
    response["Content-Disposition"] = f'attachment; filename="Salary_List_{organization_name}.pdf"'

    pisa_status = pisa.CreatePDF(html, dest=response, encoding="UTF-8")

    if pisa_status.err:
        return HttpResponse("Error generating PDF", status=500)

    return response




















# def format_date_string(date_str):
#     for fmt in ("%d %b %Y", "%d %B %Y"):  # Try both abbreviated and full month names
#         try:
#             date_obj = datetime.strptime(date_str, fmt)
#             return date_obj.strftime("%d/%m/%Y")
#         except:
#             continue
#     return date_str  # fallback if both formats fail

# Custom formatter
def format_date_string(date_str):
    for fmt in ("%Y-%m-%d", "%d-%m-%Y", "%d %b %Y", "%d %B %Y"):
        try:
            date_obj = datetime.strptime(str(date_str), fmt)
            return date_obj.strftime("%d/%m/%Y")
        except:
            continue
    return str(date_str)  # fallback as string



# def get_salary_pdf_report(organization_id, month, year):
#     with connection.cursor() as cursor:
#         cursor.execute("""
#             EXEC Employee_Payroll_SP_SalaryGenrate_PDF_Report_Select 
#                 @OrganizationID=%s, 
#                 @Month=%s, 
#                 @year=%s
#         """, [organization_id, month, year])
        
#         columns = [col[0] for col in cursor.description]
#         results = [dict(zip(columns, row)) for row in cursor.fetchall()]
#     return results




# Update status 
from django.http import JsonResponse

def update_verification(request):
    if 'OrganizationID' not in request.session:
        return redirect(MasterAttribute.Host)
    else:
        print("Show Page Session")
     
    OrganizationID = request.session["OrganizationID"]
    UserID = str(request.session["UserID"])
    if request.method == 'POST':
        employee_code = request.POST.get('employee_code')
        month = request.POST.get('month')
        year = request.POST.get('year')
        verify_type = request.POST.get('verify_type')  
        
     
     
     

        try:
            salary_slip = SalarySlip.objects.get(
                        EmployeeCode=str(employee_code),
                        month=int(month),
                        year=int(year),
                        OrganizationID=int(OrganizationID),
                        IsDelete=False
                    )

            if verify_type == 'Verify HR':
                salary_slip.HrVerify = True
            elif verify_type == 'Verify FC':
                salary_slip.FcVerify = True

       
           
           
            salary_slip.ModifyBy = UserID
            
            salary_slip.save()
            return JsonResponse({'message': 'Verification updated successfully.'})
        except SalarySlip.DoesNotExist:
            return JsonResponse({'error': 'Salary slip not found.'}, status=404)
    
    return JsonResponse({'error': 'Invalid request.'}, status=400)
# organization Config
def OrgConfigList(request):
    if 'OrganizationID' not in request.session:
        return redirect(MasterAttribute.Host)
    else:
        print("Show Page Session")
     
    OrganizationID = request.session["OrganizationID"]
    UserID = str(request.session["UserID"])
    hotelapitoken = MasterAttribute.HotelAPIkeyToken
    headers = {
        'hotel-api-token': hotelapitoken  # Replace with your actual header key and value
    }
    api_url = "http://hotelops.in/API/PyAPI/OrganizationListSelect?OrganizationID=" + str(OrganizationID)

    try:
        response = requests.get(api_url, headers=headers)
        response.raise_for_status()  # Optional: Check for any HTTP errors
        memOrg = response.json()
    except requests.exceptions.RequestException as e:
        print(f"Error occurred: {e}")   
    I = request.GET.get('I',OrganizationID)
    if I == '':
        I= OrganizationID

    
 
    orgconfig = Organization_Details.objects.filter(OID = I,IsDelete = False)
    
    context = {'orgconfig':orgconfig,'I':I,'memOrg':memOrg,'OrganizationID':OrganizationID}
    return render(request, "EMP_PAY/OrgConfig/OrgConfigList.html", context)


from pathlib import Path   
from django.urls import reverse
@transaction.atomic
def AddConfig(request):
    if 'OrganizationID' not in request.session:
        return redirect(MasterAttribute.Host)
    else:
        print("Show Page Session")
    
    OrganizationID = request.session["OrganizationID"]
    
    UserID = str(request.session["UserID"])
    hotelapitoken = MasterAttribute.HotelAPIkeyToken
    headers = {
        'hotel-api-token': hotelapitoken  # Replace with your actual header key and value
    }
    api_url = "http://hotelops.in/API/PyAPI/OrganizationListSelect?OrganizationID=" + str(OrganizationID)

    try:
        response = requests.get(api_url, headers=headers)
        response.raise_for_status()  # Optional: Check for any HTTP errors
        memOrg = response.json()
    except requests.exceptions.RequestException as e:
        print(f"Error occurred: {e}")  
    ID  = request.GET.get('ID')
    config = None
    if ID is not None:
        config = get_object_or_404(Organization_Details,id=ID,IsDelete=False)

    with transaction.atomic():
        if request.method == "POST":
             TotalMinimumDutyHours=request.POST.get('TotalMinimumDutyHours', '9') 
             TotalDutyGraceHours=request.POST.get('TotalDutyGraceHours', '0') 
             IsESICCalculate=request.POST.get('IsESICCalculate', 'Yes') 
             if ID is not None:

                OID_Code =   request.POST['OID_Code']
                OID =  request.POST['OID']
                I = OID
                
                UploadFormatType = request.POST['UploadFormatType']
                DownloadFormat = request.FILES.get('DownloadFormat')
                EndDate =request.POST['EndDate']
                
                
               
                
                config.TotalMinimumDutyHours=TotalMinimumDutyHours
                config.TotalDutyGraceHours=TotalDutyGraceHours
                config.IsESICCalculate=IsESICCalculate
                config.OID_Code =   OID_Code
                config.OID =  OID
                config.OrgUrl = request.POST['OrgUrl']
                config.cid = request.POST['cid']
                
                config.UploadFormatType = UploadFormatType



                UploadFile = DownloadFormat
                
                if UploadFile is None:
                    previous = config.DownloadFormat
                    if previous:
                        UploadFile = config.DownloadFormat.url
                        print(UploadFile)

                
                if DownloadFormat is not None:
                    if UploadFormatType != "." +Path(UploadFile).suffix.lower()[1:]:
                        messages.warning(request,f'Selected Type is not macthed with upload format')
                        
                        return redirect(reverse('AddConfig') + f'?ID={config.id}') 
                
                config.DownloadFormat = UploadFile
                config.EndDate =request.POST['EndDate']
                
                config.ModifyBy = UserID
                config.save()
                messages.success(request,"Updated Succesfully")
                 
             else:
                  
                OID_Code =   request.POST['OID_Code']
                OrgUrl = request.POST['OrgUrl']
                OID =  request.POST['OID']
                cid = request.POST['cid']
                
                I = OID
                UploadFormatType = request.POST['UploadFormatType']
                DownloadFormat = request.FILES.get('DownloadFormat')
                EndDate =request.POST['EndDate']
                
                if DownloadFormat is not None:
                    if UploadFormatType != "." +Path(DownloadFormat.name).suffix.lower()[1:]:
                        messages.warning(request,f'Selected Type is not macthed with upload format')
                        return redirect('AddConfig')
                try:
                     orgcon = Organization_Details.objects.update_or_create(
                            OID_Code =  OID_Code,
                            OID = OID,
                            OrgUrl = OrgUrl,
                            cid = cid,
                            TotalMinimumDutyHours=TotalMinimumDutyHours,
                            TotalDutyGraceHours=TotalDutyGraceHours,
                            IsESICCalculate=IsESICCalculate,
                
                            UploadFormatType = UploadFormatType,
                            DownloadFormat = DownloadFormat,
                            EndDate = EndDate,
                            
                            OrganizationID = OrganizationID,
                            CreatedBy = UserID
                        )
                     messages.warning(request,"Data is present for Selected Organization")
                except:
                        
                        messages.success(request,"Added Succesfully")
             
             return redirect(reverse('OrgConfigList') + f'?I={I}')    

    
    

    context = {'config':config,'memOrg':memOrg}
    return render(request, "EMP_PAY/OrgConfig/AddConfig.html", context)
    


@transaction.atomic
def ConfigDelete(request):
    if 'OrganizationID' not in request.session:
        return redirect(MasterAttribute.Host)
    else:
        print("Show Page Session")
     
    OrganizationID = request.session["OrganizationID"]
    UserID = str(request.session["UserID"])
    id = request.GET.get('ID')
    obj = Organization_Details.objects.get(id=id,IsDelete=False,OrganizationID = OrganizationID )
    with transaction.atomic():
       obj.IsDelete = True
       obj.ModifyBy = UserID
       I = obj.OID
       obj.save()
    return redirect(reverse('OrgConfigList') + f'?I={I}') 







# Employee weekoff Mapping


from django.shortcuts import render, redirect
from django.http import JsonResponse
import requests
from datetime import datetime
from .models import WeekOffDetails
from hotelopsmgmtpy.GlobalConfig import  MasterAttribute
from pprint import pprint

import requests
from django.shortcuts import render, redirect

def index(request):
    if 'OrganizationID' not in request.session:
        return redirect('MasterAttribute.Host')
    else:
        OrganizationID = request.session["OrganizationID"]
    
    UserID = str(request.session["UserID"])
    EmployeeCode = request.session["EmployeeCode"]
    Department_Name  =  request.session["Department_Name"]
    
    hotelapitoken = MasterAttribute.HotelAPIkeyToken  
    headers = {
        'hotel-api-token': hotelapitoken
    }
    
    # print("The Request is Reached here ----------::")
    Approval_url = "https://hotelops.in/API/PyAPI/HREmployeeListForApproval?UserID="+str(UserID)
    
    try:
        response = requests.get(Approval_url, headers=headers)
        response.raise_for_status()
        emp_list = response.json()
        emp_list.sort(key=lambda emp: emp.get('FirstName', '').lower()) 
    except requests.exceptions.RequestException as e:
        print(f"Error fetching approval list: {e}")
                
                
    today = datetime.today()
    CYear = today.year
    CMonth = today.month            
    context = {'emp_list': emp_list,'CYear':range(CYear,2020,-1),'CMonth':CMonth}
    return render(request, 'EMP_PAY/Weekoff/index.html', context)



from django.http import JsonResponse
import requests
from django.shortcuts import redirect




def all_events(request):
    if 'OrganizationID' not in request.session:
        return redirect('MasterAttribute.Host')
    
    OrganizationID = request.session["OrganizationID"]
    
    Department_Name  =  request.session["Department_Name"]
    UserID = str(request.session["UserID"])
    hotelapitoken = MasterAttribute.HotelAPIkeyToken
    headers = {
        'hotel-api-token': hotelapitoken
    }
    EmployeeSelect = request.GET.get('EmployeeSelect') or 'All'
    
    
    Approval_url = "http://hotelops.in/API/PyAPI/HREmployeeListForApproval?UserID="+str(UserID)
    
    try:
        response = requests.get(Approval_url, headers=headers)
        response.raise_for_status()
        emplist = response.json()
    except requests.exceptions.RequestException as e:
        print(f"Error fetching approval list: {e}")

    Emp_Codelist = []                
    for emp in emplist:
        
        Emp_Codelist.append(emp['EmployeeCode'])

    EmployeeDetails = []
    for emp in emplist:
           
            employee_dict = {
            "EmployeeCode": emp['EmployeeCode'],
            "EmployeeName": emp['FirstName'] + " " + emp['LastName'] 
            }
            EmployeeDetails.append(employee_dict)

        
    
    
    if not emplist:
        return JsonResponse({'error': 'No employee data found'}, status=404)
    
    
    
    if EmployeeSelect == "All":   
            all_events = WeekOffDetails.objects.filter(OrganizationID=OrganizationID, IsDelete=False, Emp_Code__in = Emp_Codelist).order_by('Emp_Code')
    else:
            all_events = WeekOffDetails.objects.filter(OrganizationID=OrganizationID, IsDelete=False, Emp_Code__in = Emp_Codelist,
                                                    Emp_Code = EmployeeSelect).order_by('Emp_Code')

    event_list = []
    
    for event in all_events:
        employee_name = next((emp['EmployeeName'] for emp in EmployeeDetails if emp['EmployeeCode'] == event.Emp_Code), '')

        event_dict = {
            'title':   employee_name   ,
            'start': event.WeekoffDate.strftime('%Y-%m-%d'),
            'end': event.WeekoffDate.strftime('%Y-%m-%d'),
            'empcode': event.Emp_Code,
            'id': event.id
        }
        event_list.append(event_dict)

    return JsonResponse(event_list, safe=False)




def add_event(request):
    if 'OrganizationID' not in request.session:
        return redirect(MasterAttribute.Host)
    else:
        OrganizationID = request.session["OrganizationID"]
    UserID = str(request.session["UserID"])
    if request.method == 'GET':
        try:
            emp_code = request.GET.get("empcode", None)
            hotelapitoken = MasterAttribute.HotelAPIkeyToken
            headers = {
                'hotel-api-token': hotelapitoken  
            }
    
            WeekOffDate = request.GET.get("WeekOffDate", None)  
            

            weekoff_id = request.GET.get("id")  
           
            if emp_code and WeekOffDate:
                
                if weekoff_id == "0":
                    existing_event = WeekOffDetails.objects.filter(WeekoffDate=WeekOffDate, OrganizationID=OrganizationID,Emp_Code=emp_code,IsDelete=False).first()
                    if existing_event:
                        return JsonResponse({'error': 'Weekoff details already exist for this employee on selected date'}, status=400)
                  
                    WeekOffDetails.objects.create(Emp_Code=emp_code, WeekoffDate=WeekOffDate, OrganizationID=OrganizationID)
                    objAtt, created = Attendance_Data.objects.update_or_create(
                        Date=WeekOffDate,
                        EmployeeCode=emp_code,
                        OrganizationID=OrganizationID,
                        defaults={
                            'IsDelete': False,  # Fields to update or set
                            # Add other fields you want to update or set here
                            'Status': "Week Off",
                        }
                    )

                    if created:
                        print("A new object was created.")
                    else:
                        print("An existing object was updated.")
                    return JsonResponse({'success': True})
                elif weekoff_id:
                    
                    try:
                        weekoff_id_int = int(weekoff_id)
                    except ValueError:
                        return JsonResponse({'error': 'Invalid weekoff ID provided'}, status=400)
                    
                    weekoff_instance = WeekOffDetails.objects.filter(id=weekoff_id_int, OrganizationID=OrganizationID).first()
                    if not weekoff_instance:
                        return JsonResponse({'error': 'Weekoff entry not found for given ID'}, status=404)
                     
                    
                    weekoff_instance.WeekoffDate = WeekOffDate
                    weekoff_instance.ModifyBy = UserID
                    weekoff_instance.save()
                    objAtt, created = Attendance_Data.objects.update_or_create(
                        Date=WeekOffDate,
                        EmployeeCode=emp_code,
                        OrganizationID=OrganizationID,
                        defaults={
                            'IsDelete': False,  # Fields to update or set
                            # Add other fields you want to update or set here
                            'Status': "Week Off",
                        }
                    )

                    if created:
                        print("A new object was created.")
                    else:
                        print("An existing object was updated.")
                    return JsonResponse({'success': True})
            else:
                return JsonResponse({'error': 'Missing parameters'}, status=400)
        except ValueError:
            return JsonResponse({'error': 'Incorrect date format. Please provide the date in MM/DD/YYYY format.'}, status=400)
        except Exception as e:
            return JsonResponse({'error': str(e)}, status=500)



def remove(request):
    if 'OrganizationID' not in request.session:
        return redirect(MasterAttribute.Host)
    else:
        OrganizationID = request.session["OrganizationID"]
    UserID = str(request.session["UserID"])
    if request.method == 'GET':
        id = request.GET.get("id", None)
        if id:
            obj = WeekOffDetails.objects.get(id=id)
            obj.IsDelete = True
            obj.ModifyBy = UserID
            obj.save()

            return JsonResponse({'success': True})
        else:
            return JsonResponse({'error': 'Missing ID'}, status=400)






import requests
from datetime import datetime, timedelta
from openpyxl import Workbook
from django.http import HttpResponse
import io

def download_weekoff(request):
    if 'OrganizationID' not in request.session:
        return redirect('MasterAttribute.Host')

    OrganizationID = request.session["OrganizationID"]
    
    UserID = str(request.session["UserID"])
    hotelapitoken = MasterAttribute.HotelAPIkeyToken

    headers = {
        'hotel-api-token': hotelapitoken
    }

    Approval_url = f"http://hotelops.in/API/PyAPI/HREmployeeListForApproval?UserID={UserID}"

    try:
        response = requests.get(Approval_url, headers=headers)
        response.raise_for_status()
        emplist = response.json()
    except requests.exceptions.RequestException as e:
        print(f"Error fetching approval list: {e}")
        return HttpResponse("Error fetching approval list.", status=500)

    year = int(request.GET.get('year'))
    month = int(request.GET.get('month'))

    wb = Workbook()
    ws = wb.active
    ws.title = "Weekoffs"
    ws.append(["EmployeeName" ,"Department","EmployeeCode", "WeekoffDate"])

    start_date = datetime(year, month, 1)
    while start_date.weekday() != 6:  
        start_date += timedelta(days=1)

    sundays = []
    while start_date.month == month:
        sundays.append(start_date.strftime("%d-%m-%Y"))
        start_date += timedelta(days=7)

    for emp in emplist:
            employee_name = f"{emp['FirstName']} {emp['MiddleName']} {emp['LastName']}".strip()
            department = emp['Department']
            employee_code = emp['EmployeeCode']
            for sunday in sundays:
                ws.append([employee_name, department, employee_code, sunday])


    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    response = HttpResponse(output, content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = f'attachment; filename=Weekoffs_{year}_{month}.xlsx'
    return response


from datetime import datetime
import pandas as pd
from django.http import JsonResponse
from django.shortcuts import redirect
from .models import WeekOffDetails
import os

def upload_weekoff(request):
    if 'OrganizationID' not in request.session:
        print("OrganizationID not found in session.")
        return redirect(MasterAttribute.Host)
    
    OrganizationID = request.session.get("OrganizationID")
    UserID = str(request.session.get("UserID"))
    
    if request.method == 'POST':
        error_messages = []

        if 'file' not in request.FILES:
            return JsonResponse({"error": "File not found in request."})

        excel_file = request.FILES['file']
        filename = excel_file.name.lower()

        try:
            if filename.endswith(".xlsx"):
                df = pd.read_excel(excel_file, engine="openpyxl")
            elif filename.endswith(".xls"):
                df = pd.read_excel(excel_file, engine="xlrd")
            elif filename.endswith(".csv"):
                df = pd.read_csv(excel_file)
            else:
                return JsonResponse({"error": "Unsupported file format. Please upload .xlsx, .xls, or .csv"})

            # print("df columns:", df.columns)
            
            for index, row in df.iterrows():
                emp_code = row.get('EmployeeCode')
                weekdate = row.get('WeekoffDate')
                
                if not emp_code or not weekdate:
                    error_message = f"Row {index + 1}: Employee code or week-off date missing."
                    # print(error_message)
                    # error_messages.append(error_message)
                    continue
                
                try:
                    parsed_date = datetime.strptime(weekdate, '%d-%m-%Y')
                    # weekoff_date = weekdate.strftime('%Y-%m-%d')
                except ValueError as ve:
                    try:
                        parsed_date = datetime.strptime(weekdate, '%Y-%m-%d')
                        # weekoff_date = weekdate.strftime('%Y-%m-%d')
                    except ValueError as ve:
                        error_message = f"Row {index + 1}: Date format error: {ve}"
                        # print(error_message)
                        # error_messages.append(error_message)
                        continue
                
                try:
                    existing_event = WeekOffDetails.objects.filter(
                        WeekoffDate=parsed_date,
                        OrganizationID=OrganizationID,
                        Emp_Code=emp_code,
                        IsDelete=False
                    ).first()
                    
                    if existing_event:
                        error_message = f"Row {index + 1}: Weekoff details already exist for employee {emp_code} on {weekoff_date}."
                        # print(error_message)
                        # error_messages.append(error_message)
                        continue
                    
                    WeekOffDetails.objects.create(
                        Emp_Code=emp_code,
                        WeekoffDate=parsed_date,
                        OrganizationID=OrganizationID,
                        CreatedBy=UserID
                    )
                    objAtt, created = Attendance_Data.objects.update_or_create(
                        Date=parsed_date,
                        EmployeeCode=emp_code,
                        OrganizationID=OrganizationID,
                        defaults={
                            'IsDelete': False,  # Fields to update or set
                            # Add other fields you want to update or set here
                            'Status': "Week Off",
                        }
                    )

                    if created:
                        print("A new object was created.")
                    else:
                        print("An existing object was updated.")
                    
                except Exception as e:
                    error_message = f"Row {index + 1}: An error occurred while processing: {e}"
                    # print(error_message)
                    # error_messages.append(error_message)
                    continue
            
            if error_messages:
                return JsonResponse({"error": "Some rows contained errors.", "details": error_messages})
            
            print("Weekoff details uploaded successfully.")
            return JsonResponse({"success": "Weekoff details uploaded successfully"})
        
        except Exception as e:
            print(f"An error occurred: {e}")
            return JsonResponse({"error": f"An error occurred: {e}"})
            
    print("Invalid request method.")
    return JsonResponse({"error": "Invalid request method."})

    


# Shfit Managerer
def ShfitList(request):
    if 'OrganizationID' not in request.session:
        return redirect(MasterAttribute.Host)
    else:
        print("Show Page Session")
     
    OrganizationID = request.session["OrganizationID"]
    UserID = str(request.session["UserID"])
    hotelapitoken = MasterAttribute.HotelAPIkeyToken
    headers = {
        'hotel-api-token': hotelapitoken
    }
    
    Approval_url = "http://hotelops.in/API/PyAPI/HREmployeeListForApproval?UserID="+str(UserID)
    
    try:
        response = requests.get(Approval_url, headers=headers)
        response.raise_for_status()
        emp_list = response.json()
        emp_list.sort(key=lambda emp: emp.get('EmployeeCode', '').lower()) 
    except requests.exceptions.RequestException as e:
        print(f"Error fetching approval list: {e}")
        emp_list = []  
    
   
    filtered_emp_list = []
    for emp in emp_list:
        EmployeeCode = emp['EmployeeCode']
        shfit =  ShfitMaster.objects.filter(EmployeeCode=EmployeeCode, OrganizationID=OrganizationID, IsDelete=False).first()
        if shfit:
            emp['ShfitType'] = shfit.ShfitType
            filtered_emp_list.append(emp)
    
    context = {"emp_list": filtered_emp_list}
    return render(request, "EMP_PAY/Shfit/ShfitList.html", context)



from django.http import JsonResponse

def UpdateShift(request):
    if 'OrganizationID' not in request.session:
        return redirect(MasterAttribute.Host)
    else:
        print("Show Page Session")
     
    OrganizationID = request.session["OrganizationID"]
    UserID = str(request.session["UserID"])
    if request.method == "GET":
        EmployeeCode = request.GET.get('EmployeeCode')
        ShiftType = request.GET.get('ShfitType')  

        obj = get_object_or_404(ShfitMaster,EmployeeCode=EmployeeCode, OrganizationID = OrganizationID, IsDelete=False)
       
        obj.ShfitType = ShiftType  
        obj.ModifyBy = UserID
        obj.save()

        response_data = {
            'message': f'Status {ShiftType} updated successfully'  
        }
        return JsonResponse(response_data, status=200)



def IntialShfitCreate(request):
    if 'OrganizationID' not in request.session:
        return redirect(MasterAttribute.Host)
    else:
        print("Show Page Session")
     
    OrganizationID = request.session["OrganizationID"]
    UserID = str(request.session["UserID"])
    hotelapitoken = MasterAttribute.HotelAPIkeyToken
    headers = {
        'hotel-api-token': hotelapitoken
    }
    
    Approval_url = "http://hotelops.in/API/PyAPI/HREmployeeListForApproval?UserID="+str(UserID)
    
    try:
        response = requests.get(Approval_url, headers=headers)
        response.raise_for_status()
        emp_list = response.json()
        emp_list.sort(key=lambda emp: emp.get('EmployeeCode', '').lower()) 
    except requests.exceptions.RequestException as e:
        print(f"Error fetching approval list: {e}")

    create  = request.GET.get('create')
    if create == "True":
        for emp in emp_list:
            EmployeeCode = emp['EmployeeCode']
            ShfitDetails =  ShfitMaster.objects.filter(EmployeeCode= EmployeeCode,OrganizationID =OrganizationID,IsDelete=False).first()
            if ShfitDetails:
                messages.warning(request,f'Shfit is present for {EmployeeCode}')
            ShfitMaster.objects.create(EmployeeCode= EmployeeCode,OrganizationID =OrganizationID,CreatedBy = UserID)    



    messages.success(request,"Created Successfully")
    return redirect('ShfitList')





from django.db.models import Max, Min
from django.shortcuts import redirect, render
from django.contrib import messages
from .models import Raw_Attendance_Data, WeekOffDetails, Attendance_Data
from datetime import datetime, timedelta
import requests

@transaction.atomic
def DumpAttendace(request):
    if 'OrganizationID' not in request.session:
        return redirect(MasterAttribute.Host)
    
    OrganizationID = request.session["OrganizationID"]
    UserID = str(request.session["UserID"])
    
    hotelapitoken = MasterAttribute.HotelAPIkeyToken
    # headers = {
    #     'hotel-api-token': hotelapitoken  
    # }
    # api_url = "https://hotelops.in/API/PyAPI/HREmployeeList?OrganizationID="+str(OrganizationID)
    
    # try:
    #     response = requests.get(api_url, headers=headers)
    #     response.raise_for_status() 
    #     emp_list = response.json()
    # except requests.exceptions.RequestException as e:
    #     print(f"Error occurred: {e}")
    emp_list =  EmployeeDataSelect(OrganizationID)    
    if request.method == "POST":
        S_Date_str = request.POST['S_Date']
        E_Date_str = request.POST['E_Date']

        S_Date = datetime.strptime(S_Date_str, '%Y-%m-%d').date()
        E_Date = datetime.strptime(E_Date_str, '%Y-%m-%d').date()

        date_range = [S_Date + timedelta(days=x) for x in range((E_Date - S_Date).days + 1)]
        print(date_range)
        for attendance_date_obj in date_range:
            for emp in emp_list:
                EmployeeCode = emp['EmployeeCode']

                latest_out_time_subquery = Raw_Attendance_Data.objects.filter(
                    EmployeeCode=EmployeeCode,
                    Date=attendance_date_obj,
                    OrganizationID = OrganizationID  ,
                    IsDelete=False,
                    Status='OUT'
                ).aggregate(Max('Time'))

                earliest_in_time_subquery = Raw_Attendance_Data.objects.filter(
                    EmployeeCode=EmployeeCode,
                     OrganizationID = OrganizationID  ,
                    IsDelete=False,
                    Date=attendance_date_obj,
                    Status='IN'
                ).aggregate(Min('Time'))

                latest_out_time = latest_out_time_subquery['Time__max']
                earliest_in_time = earliest_in_time_subquery['Time__min']

                if latest_out_time and earliest_in_time:
                    in_time = datetime.strptime(earliest_in_time, '%H:%M:%S')
                    out_time = datetime.strptime(latest_out_time, '%H:%M:%S')
                    duty_hours = out_time - in_time
                    duty_hours_time = str(duty_hours).split()[-1]

                    attendance_records = {
                        'EmployeeCode': EmployeeCode,
                        'Date': attendance_date_obj,
                        'MinInTime': earliest_in_time,
                        'MaxOutTime': latest_out_time,
                        'DutyHours': duty_hours_time,
                    }

                    status = "Absent"
                    if duty_hours >= timedelta(hours=9):
                        status = "Present"
                    elif duty_hours >= timedelta(hours=5):
                        status = "Half Day Present"

                    attendance_records['status'] = status

                    Attendance_Data.objects.create(
                        EmployeeCode=EmployeeCode,
                        Date=attendance_date_obj,
                        In_Time=attendance_records['MinInTime'],
                        Out_Time=attendance_records['MaxOutTime'],
                        Duty_Hour=attendance_records['DutyHours'],
                        OrganizationID=OrganizationID,
                        CreatedBy=UserID,
                        Status=attendance_records['status'],
                        IsUpload=True
                    )

                else:
                    try:
                        leave = Leave_Application.objects.get(
                            Start_Date__lte=attendance_date_obj,
                            End_Date__gte=attendance_date_obj,
                            IsDelete=False,
                            Emp_code=EmployeeCode,
                            Status=1,
                             OrganizationID = OrganizationID  ,
                   
                        )
                        status = leave.Leave_Type_Master.Type
                    except Leave_Application.DoesNotExist:
                        try:
                            weekoff = WeekOffDetails.objects.get(
                                Emp_Code=EmployeeCode,
                                WeekoffDate=attendance_date_obj,
                                IsDelete=False,
                                 OrganizationID = OrganizationID  ,
                  
                            )
                            status = "WeekOff"
                        except WeekOffDetails.DoesNotExist:
                            status = "Absent"

                    Attendance_Data.objects.create(
                        EmployeeCode=EmployeeCode,
                        Date=attendance_date_obj,
                        In_Time=None,
                        Out_Time=None,
                        Duty_Hour=None,
                        OrganizationID=OrganizationID,
                        CreatedBy=UserID,
                        Status=status,
                        IsUpload=True
                    )

        messages.success(request, 'Attendance data dumped successfully.')
        return redirect('DumpAttendace')
    
    context = {}    
    return render(request, "EMP_PAY/Attendance/DumpAttendace.html", context)



def parse_date(raw_value):
    if not raw_value or raw_value.strip() == '':
        return None

    value = raw_value.replace('="', '').replace('"', '').strip()
    # value = raw_value.replace('="', '').replace('"', '').strip().title()

    date_formats = [
        "%d-%b-%y",  
        "%d/%m/%y",   
        "%d/%m/%Y",   
        "%Y-%m-%d",   
        "%d-%m-%Y",   
    ]

    for fmt in date_formats:
        try:
            return datetime.strptime(value, fmt).date()
        except ValueError:
            continue
    return None



def Alifupload_csv(request):
    if 'OrganizationID' not in request.session:
        return redirect(MasterAttribute.Host)
    
    context = {}
    # return render(request, "EMP_PAY/Attendance/Alifupload_csv.html", {'form':form})
    return render(request, "EMP_PAY/Attendance/Alifupload_csv.html", context)
    


# -----------------------------------------------------

from rest_framework.views import APIView
from rest_framework.response import Response
from rest_framework import status
from django.db import transaction, connection
import csv, io
from datetime import datetime
from .serializers import AlifFileUploadSerializer

class AlifUploadCSVApi(APIView):

    def post(self, request, *args, **kwargs):
        OrganizationID = request.session.get("OrganizationID")
        UserID = str(request.session.get("UserID"))

        if not OrganizationID:
            return Response({"error": "Unauthorized: OrganizationID not found in session."},
                            status=status.HTTP_401_UNAUTHORIZED)

        serializer = AlifFileUploadSerializer(data=request.data)
        if not serializer.is_valid():
            return Response(serializer.errors, status=status.HTTP_400_BAD_REQUEST)

        file = serializer.validated_data['file']
        csv_reader = csv.reader(io.TextIOWrapper(file.file, encoding='utf-8'))

        # Skip header rows
        next(csv_reader, None)
        next(csv_reader, None)

        all_punch_dates = []
        errors = []
        new_records = []

        for row_index, row in enumerate(csv_reader, start=1):
            try:
                punch_date = parse_date(row[4])
                if punch_date:
                    all_punch_dates.append(punch_date)

                    emp_code = row[0].replace('="', '').replace('"', '') if row[0].strip() else None
                    employee_name = row[1] if row[1].strip() else None
                    punch_limit = int(row[2].replace('="', '').replace('"', '')) if row[2].strip() else 0
                    working_hours = row[3].replace('="', '').replace('"', '') if row[3].strip() else None
                    attend_code = row[5] if row[5].strip() else None
                    first_in_punch = datetime.strptime(row[6].replace('="', '').replace('"', ''), "%H.%M").time() if row[6].strip() else None
                    last_out_punch = datetime.strptime(row[7].replace('="', '').replace('"', ''), "%H.%M").time() if row[7].strip() else None
                    worked_hours = row[8].replace('="', '').replace('"', '') if row[8].strip() else None
                    opr_id = row[9] if row[9].strip() else None
                    opr_date = datetime.strptime(row[10].replace('="', '').replace('"', ''), "%d-%b-%y").date() if row[10].strip() else None

                    new_records.append(
                        AlifCSVPunchRecord(
                            emp_code=emp_code,
                            employee_name=employee_name,
                            punch_limit=punch_limit,
                            working_hours=working_hours,
                            punch_date=punch_date,
                            attend_code=attend_code,
                            first_in_punch=first_in_punch,
                            last_out_punch=last_out_punch,
                            worked_hours=worked_hours,
                            opr_id=opr_id,
                            opr_date=opr_date,
                            OrganizationID=OrganizationID,
                            CreatedBy=UserID
                        )
                    )
            except Exception as e:
                errors.append(f"Row {row_index}: {str(e)}")

        if not all_punch_dates:
            return Response({"error": "No valid punch dates found in the file."}, status=status.HTTP_400_BAD_REQUEST)

        first_date, last_date = all_punch_dates[0], all_punch_dates[-1]

        # Backup and replace old records
        backup_and_replace_punch_records(OrganizationID, first_date, last_date)

        # Insert new records
        if new_records:
            AlifCSVPunchRecord.objects.bulk_create(new_records)

        # print("stored procedure is called now...")
        # Call stored procedure
        with transaction.atomic():
            with connection.cursor() as cursor:
                cursor.execute("EXEC Employee_Payroll_Alif_UpdateAttendance_Org")

        response_data = {
            "inserted_records": len(new_records),
            "errors": errors,
            "status": "File uploaded and processed successfully" if not errors else "Completed with errors"
        }
        return Response(response_data, status=status.HTTP_200_OK)

# -----------------------------------------------------

def backup_and_replace_punch_records(org_id, first_date, last_date):

    with transaction.atomic():
        old_records = AlifCSVPunchRecord.objects.filter(
            OrganizationID=org_id,
            punch_date__range=(first_date, last_date)
        )

        # print("Old Records Count:", old_records.count())

        if old_records.exists():
            # print("Backing up old records...")
            log_entries = [
                AlifCSVPunchRecord_Log(
                    **{field.name: getattr(record, field.name)
                       for field in AlifCSVPunchRecord._meta.fields
                       if field.name != "id"}  # exclude primary key if needed
                )
                for record in old_records
            ]
            AlifCSVPunchRecord_Log.objects.bulk_create(log_entries)
            # print(f"Backed up {len(log_entries)} records to AlifCSVPunchRecord_Log.")

            # print(f"Deleted {old_records.count()} old records from AlifCSVPunchRecord.")
            old_records.delete()
        else:           
            print("No old records to back up.")

      